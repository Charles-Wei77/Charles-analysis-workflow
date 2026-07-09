#!/usr/bin/env python3
from __future__ import annotations

import datetime as dt
import json
import math
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
WORKFLOWS_ROOT = ROOT.parent
HUB = Path(os.getenv("AGENT_SKILLS_HUB", str(WORKFLOWS_ROOT.parent))).expanduser()
RECALC = Path(
    os.getenv(
        "REPORT_WORKFLOW_RECALC_SCRIPT",
        str(HUB / "skills" / "github-skill-anthropics-skills-skills-xlsx" / "scripts" / "recalc.py"),
    )
).expanduser()
WORKFLOW_PYTHON = Path(os.getenv("WORKFLOW_PYTHON", sys.executable)).expanduser()


BLUE = "1F4E79"
LIGHT_BLUE = "D9E1F2"
MED_BLUE = "BDD7EE"
LIGHT_GREY = "F2F2F2"
WHITE = "FFFFFF"
FONT_BLUE = "0000FF"
FONT_GREEN = "008000"
FONT_BLACK = "000000"


def safe_float(value: Any, default: float = 0.0) -> float:
    try:
        if value is None or value == "":
            return default
        return float(value)
    except Exception:
        return default


def pct(value: float) -> float:
    return value / 100 if abs(value) > 1 else value


def slug(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9\u4e00-\u9fff]+", "-", text).strip("-")[:80] or "model"


def source_title(src: Any) -> str:
    if isinstance(src, dict):
        return src.get("title") or src.get("url") or "Source"
    return getattr(src, "title", "Source")


def source_url(src: Any) -> str:
    if isinstance(src, dict):
        return src.get("url", "")
    return getattr(src, "url", "")


def source_date(src: Any) -> str:
    if isinstance(src, dict):
        return src.get("date", "")
    return getattr(src, "date", "")


def source_evidence(src: Any) -> str:
    if isinstance(src, dict):
        return src.get("evidence") or src.get("note") or ""
    return getattr(src, "evidence", "") or getattr(src, "note", "")


def default_historical() -> list[dict[str, Any]]:
    return [
        {"year": 2023, "revenue": 4200.0, "gross_margin": 0.295, "ebit_margin": 0.035, "net_income": 130.0, "da": 130.0, "capex": 170.0, "nwc": 620.0, "cash": 900.0, "debt": 1450.0, "shares": 870.0, "total_assets": 9600.0, "equity": 3000.0, "source": "Placeholder historicals"},
        {"year": 2024, "revenue": 4000.0, "gross_margin": 0.290, "ebit_margin": -0.140, "net_income": -810.0, "da": 150.0, "capex": 210.0, "nwc": 700.0, "cash": 780.0, "debt": 1500.0, "shares": 870.0, "total_assets": 10100.0, "equity": 2300.0, "source": "Placeholder historicals"},
        {"year": 2025, "revenue": 4888.0, "gross_margin": 0.2945, "ebit_margin": 0.035, "net_income": 45.0, "da": 160.0, "capex": 230.0, "nwc": 650.0, "cash": 950.0, "debt": 1450.0, "shares": 871.0, "total_assets": 9415.0, "equity": 2450.0, "source": "Placeholder historicals"},
    ]


def normalize_model_data(data: dict[str, Any], sources: list[Any]) -> dict[str, Any]:
    fm = data.get("financial_model") or data.get("financials") or {}
    historical = fm.get("historical") or data.get("historical_financials") or default_historical()
    normalized = []
    for item in historical:
        year = int(item.get("year"))
        revenue = safe_float(item.get("revenue"))
        gross_margin = pct(safe_float(item.get("gross_margin"), 0.30))
        gross_profit = safe_float(item.get("gross_profit"), revenue * gross_margin)
        ebit_margin = pct(safe_float(item.get("ebit_margin"), safe_float(item.get("ebit"), revenue * 0.05) / revenue if revenue else 0.05))
        ebit = safe_float(item.get("ebit"), revenue * ebit_margin)
        net_income = safe_float(item.get("net_income"), ebit * 0.75)
        da = safe_float(item.get("da") or item.get("depreciation_amortization"), revenue * 0.035)
        capex = safe_float(item.get("capex"), revenue * 0.045)
        nwc = safe_float(item.get("nwc") or item.get("working_capital"), revenue * 0.13)
        cash = safe_float(item.get("cash"), revenue * 0.18)
        debt = safe_float(item.get("debt"), revenue * 0.28)
        shares = safe_float(item.get("shares") or item.get("shares_m"), 900.0)
        total_assets = safe_float(item.get("total_assets"), cash + nwc + revenue * 0.65 + revenue * 0.45)
        equity = safe_float(item.get("equity"), total_assets - debt - revenue * 0.55)
        normalized.append({
            "year": year,
            "revenue": revenue,
            "gross_profit": gross_profit,
            "gross_margin": gross_profit / revenue if revenue else gross_margin,
            "ebit": ebit,
            "ebit_margin": ebit / revenue if revenue else ebit_margin,
            "net_income": net_income,
            "da": da,
            "capex": capex,
            "nwc": nwc,
            "cash": cash,
            "debt": debt,
            "shares": shares,
            "total_assets": total_assets,
            "equity": equity,
            "source": item.get("source") or "User/model input",
        })
    normalized = sorted(normalized, key=lambda x: x["year"])[-3:]
    latest = normalized[-1]
    segments = fm.get("segments") or data.get("segments")
    if not segments:
        segments = [
            {"name": "工业机器人及智能制造系统", "revenue": latest["revenue"] * 0.82, "gross_margin": 0.292},
            {"name": "自动化核心部件及运动控制系统", "revenue": latest["revenue"] * 0.18, "gross_margin": 0.304},
        ]
    segments = [{"name": s.get("name", f"Segment {i+1}"), "revenue": safe_float(s.get("revenue"), latest["revenue"] / max(1, len(segments))), "gross_margin": pct(safe_float(s.get("gross_margin"), 0.30))} for i, s in enumerate(segments)]
    market = {
        "price": safe_float((fm.get("market") or {}).get("price") or data.get("price"), 10.0),
        "shares": safe_float((fm.get("market") or {}).get("shares") or (fm.get("market") or {}).get("shares_m") or latest["shares"], latest["shares"]),
        "cash": safe_float((fm.get("market") or {}).get("cash"), latest["cash"]),
        "debt": safe_float((fm.get("market") or {}).get("debt"), latest["debt"]),
        "risk_free_rate": pct(safe_float((fm.get("market") or {}).get("risk_free_rate"), 0.022)),
        "beta": safe_float((fm.get("market") or {}).get("beta"), 1.20),
        "equity_risk_premium": pct(safe_float((fm.get("market") or {}).get("equity_risk_premium"), 0.055)),
        "pretax_cost_of_debt": pct(safe_float((fm.get("market") or {}).get("pretax_cost_of_debt"), 0.045)),
        "tax_rate": pct(safe_float((fm.get("market") or {}).get("tax_rate"), 0.25)),
    }
    scenarios = fm.get("scenarios") or default_scenarios(len(segments))
    comps = fm.get("comps") or [
        {"company": "汇川技术", "ticker": "300124.SZ", "market_cap": 0, "ev": 0, "revenue": 0, "net_income": 0, "source": "待接入用户自配结构化数据源"},
        {"company": "埃夫特", "ticker": "688165.SH", "market_cap": 0, "ev": 0, "revenue": 0, "net_income": 0, "source": "待接入用户自配结构化数据源"},
        {"company": "拓斯达", "ticker": "300607.SZ", "market_cap": 0, "ev": 0, "revenue": 0, "net_income": 0, "source": "待接入用户自配结构化数据源"},
    ]
    return {
        "company": data.get("company") or "Company",
        "stock_code": data.get("stock_code") or data.get("code") or "",
        "industry": data.get("industry") or "",
        "topic": data.get("topic") or "",
        "currency": fm.get("currency", "RMB"),
        "units": fm.get("units", "RMB mm"),
        "historical": normalized,
        "segments": segments,
        "market": market,
        "scenarios": scenarios,
        "comps": comps,
        "sources": sources,
        "data_policy": fm.get("data_policy", "Primary: official filings and exchange disclosures; structured fields come from user-configured sources such as Wudao MCP or private licensed providers."),
    }


def default_scenarios(segment_count: int) -> dict[str, dict[str, Any]]:
    base_segment_growth = {
        "bear": [-0.02, 0.03, 0.04, 0.04, 0.03],
        "base": [0.08, 0.10, 0.10, 0.09, 0.08],
        "bull": [0.15, 0.14, 0.13, 0.11, 0.10],
    }
    scenarios: dict[str, dict[str, Any]] = {}
    for name in ["bear", "base", "bull"]:
        scenarios[name] = {
            "segment_growth": [base_segment_growth[name][:] for _ in range(segment_count)],
            "gross_margin": {"bear": [0.285, 0.286, 0.287, 0.288, 0.289], "base": [0.300, 0.305, 0.310, 0.315, 0.318], "bull": [0.315, 0.322, 0.328, 0.333, 0.338]}[name],
            "ebit_margin": {"bear": [0.025, 0.030, 0.035, 0.040, 0.045], "base": [0.050, 0.060, 0.070, 0.080, 0.090], "bull": [0.075, 0.090, 0.105, 0.115, 0.125]}[name],
            "tax_rate": [0.25] * 5,
            "da_pct_revenue": [0.035] * 5,
            "capex_pct_revenue": {"bear": [0.055] * 5, "base": [0.050] * 5, "bull": [0.047] * 5}[name],
            "nwc_pct_delta_revenue": {"bear": [0.080] * 5, "base": [0.060] * 5, "bull": [0.050] * 5}[name],
            "terminal_growth": {"bear": 0.020, "base": 0.025, "bull": 0.030}[name],
            "wacc": {"bear": 0.115, "base": 0.100, "bull": 0.090}[name],
        }
    return scenarios


class ModelBuilder:
    def __init__(self, model: dict[str, Any], out_dir: Path):
        self.model = model
        self.out_dir = out_dir
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.projection_years = [self.model["historical"][-1]["year"] + i for i in range(1, 6)]
        self.hist_years = [x["year"] for x in self.model["historical"]]
        self.all_years = self.hist_years + self.projection_years
        self.rows: dict[str, dict[str, int]] = {}
        self.scenario_rows: dict[tuple[str, str], int] = {}

    def build(self) -> dict[str, Any]:
        self.build_cover()
        self.build_sources()
        self.build_historical()
        self.build_revenue_build()
        self.build_income_statement()
        self.build_balance_sheet()
        self.build_cash_flow()
        self.build_wacc()
        self.build_dcf()
        self.build_comps()
        self.build_sensitivity()
        self.build_charts()
        self.build_audit_sheet()
        self.format_workbook()
        self.wb.calculation.calcMode = "auto"
        self.wb.calculation.fullCalcOnLoad = True
        self.wb.calculation.forceFullCalc = True
        model_path = self.out_dir / "financial_model.xlsx"
        self.wb.save(model_path)
        valuation = self.compute_valuation_summary()
        valuation_path = self.out_dir / "valuation_summary.json"
        valuation_path.write_text(json.dumps(valuation, ensure_ascii=False, indent=2), encoding="utf-8")
        recalc = self.run_recalc(model_path)
        audit_path = self.out_dir / "model_audit.md"
        audit_path.write_text(self.audit_markdown(model_path, valuation, recalc), encoding="utf-8")
        return {
            "model_path": str(model_path),
            "valuation_path": str(valuation_path),
            "audit_path": str(audit_path),
            "valuation_summary": valuation,
            "recalc": recalc,
        }

    def ws(self, name: str):
        return self.wb.create_sheet(name)

    def input(self, ws, cell: str, value: Any, comment: str = ""):
        c = ws[cell]
        c.value = value
        c.font = Font(color=FONT_BLUE)
        c.fill = PatternFill("solid", fgColor=LIGHT_GREY)
        if comment:
            c.comment = Comment(f"Source: {comment}", "Codex")
        return c

    def formula(self, ws, cell: str, formula: str, link: bool = False):
        c = ws[cell]
        c.value = formula if formula.startswith("=") else f"={formula}"
        c.font = Font(color=FONT_GREEN if link else FONT_BLACK)
        return c

    def section(self, ws, row: int, title: str, end_col: int = 8):
        ws.cell(row, 1).value = title
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_col)
        cell = ws.cell(row, 1)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="left")

    def header_row(self, ws, row: int, labels: list[Any], start_col: int = 1):
        for i, label in enumerate(labels, start_col):
            c = ws.cell(row, i)
            c.value = label
            c.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center")

    def build_cover(self):
        ws = self.ws("Cover")
        self.section(ws, 1, f"{self.model['company']} Financial Model", 6)
        rows = [
            ("Company", self.model["company"]),
            ("Ticker", self.model["stock_code"]),
            ("Industry", self.model["industry"]),
            ("Topic", self.model["topic"]),
            ("Model depth", "Heavy banking: segment revenue build + 3-statement + WACC + DCF + comps + audit"),
            ("Data policy", self.model["data_policy"]),
            ("Generated at", dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ]
        for r, (label, value) in enumerate(rows, 3):
            ws.cell(r, 1).value = label
            self.input(ws, f"B{r}", value, "Workflow metadata")

    def build_sources(self):
        ws = self.ws("Sources")
        self.section(ws, 1, "Source Register", 6)
        self.header_row(ws, 3, ["ID", "Title", "Date", "URL", "Evidence"])
        for i, src in enumerate(self.model["sources"], 4):
            ws.cell(i, 1).value = f"S{i-3}"
            ws.cell(i, 2).value = source_title(src)
            ws.cell(i, 3).value = source_date(src)
            ws.cell(i, 4).value = source_url(src)
            ws.cell(i, 5).value = source_evidence(src)[:500]

    def build_historical(self):
        ws = self.ws("Historical Financials")
        self.section(ws, 1, f"Historical Financials ({self.model['units']})", 6)
        self.header_row(ws, 3, ["Line item"] + [f"FY{y}A" for y in self.hist_years])
        lines = [
            ("Revenue", "revenue"),
            ("Gross Profit", "gross_profit"),
            ("Gross Margin", "gross_margin"),
            ("EBIT", "ebit"),
            ("EBIT Margin", "ebit_margin"),
            ("Net Income", "net_income"),
            ("D&A", "da"),
            ("CapEx", "capex"),
            ("Net Working Capital", "nwc"),
            ("Cash", "cash"),
            ("Debt", "debt"),
            ("Diluted Shares (M)", "shares"),
            ("Total Assets", "total_assets"),
            ("Equity", "equity"),
        ]
        self.rows["Historical Financials"] = {}
        for r, (label, key) in enumerate(lines, 4):
            ws.cell(r, 1).value = label
            self.rows["Historical Financials"][key] = r
            for col, item in enumerate(self.model["historical"], 2):
                self.input(ws, f"{get_column_letter(col)}{r}", item[key], f"{item.get('source', 'Historical financial input')}, FY{item['year']}")

    def build_revenue_build(self):
        ws = self.ws("Revenue Build")
        self.section(ws, 1, "Segment Revenue Build And Scenario Assumptions", 8)
        self.input(ws, "B4", 2, "Case selector: 1=Bear, 2=Base, 3=Bull")
        ws["A4"] = "Case Selector"
        self.formula(ws, "B5", '=IF(B4=1,"Bear",IF(B4=2,"Base","Bull"))')
        ws["A5"] = "Selected Case"
        col_labels = ["Assumption"] + [f"FY{y}E" for y in self.projection_years]
        row = 7
        scenario_rows: dict[tuple[str, str], int] = {}
        for scenario_idx, scenario in enumerate(["bear", "base", "bull"], 1):
            self.section(ws, row, f"{scenario.upper()} CASE ASSUMPTIONS", 7)
            self.header_row(ws, row + 1, col_labels)
            row += 2
            for si, segment in enumerate(self.model["segments"]):
                ws.cell(row, 1).value = f"Growth - {segment['name']}"
                scenario_rows[(scenario, f"segment_{si}")] = row
                values = self.model["scenarios"][scenario]["segment_growth"][si]
                for j, value in enumerate(values, 2):
                    self.input(ws, f"{get_column_letter(j)}{row}", value, f"{scenario} scenario assumption")
                row += 1
            for key, label in [
                ("gross_margin", "Gross Margin"),
                ("ebit_margin", "EBIT Margin"),
                ("tax_rate", "Tax Rate"),
                ("da_pct_revenue", "D&A % Revenue"),
                ("capex_pct_revenue", "CapEx % Revenue"),
                ("nwc_pct_delta_revenue", "NWC % Delta Revenue"),
            ]:
                ws.cell(row, 1).value = label
                scenario_rows[(scenario, key)] = row
                for j, value in enumerate(self.model["scenarios"][scenario][key], 2):
                    self.input(ws, f"{get_column_letter(j)}{row}", value, f"{scenario} scenario assumption")
                row += 1
            for key, label in [("terminal_growth", "Terminal Growth"), ("wacc", "WACC")]:
                ws.cell(row, 1).value = label
                scenario_rows[(scenario, key)] = row
                for j in range(2, 7):
                    self.input(ws, f"{get_column_letter(j)}{row}", self.model["scenarios"][scenario][key], f"{scenario} scenario assumption")
                row += 1
            row += 1

        self.section(ws, row, "SELECTED CASE ASSUMPTIONS", 7)
        self.header_row(ws, row + 1, col_labels)
        selected_start = row + 2
        selected_rows: dict[str, int] = {}
        row = selected_start
        for si, segment in enumerate(self.model["segments"]):
            key = f"segment_{si}"
            ws.cell(row, 1).value = f"Selected Growth - {segment['name']}"
            selected_rows[key] = row
            for j in range(2, 7):
                rows = [scenario_rows[(s, key)] for s in ["bear", "base", "bull"]]
                col = get_column_letter(j)
                self.formula(ws, f"{col}{row}", f"=CHOOSE($B$4,{col}{rows[0]},{col}{rows[1]},{col}{rows[2]})")
            row += 1
        for key, label in [
            ("gross_margin", "Selected Gross Margin"),
            ("ebit_margin", "Selected EBIT Margin"),
            ("tax_rate", "Selected Tax Rate"),
            ("da_pct_revenue", "Selected D&A % Revenue"),
            ("capex_pct_revenue", "Selected CapEx % Revenue"),
            ("nwc_pct_delta_revenue", "Selected NWC % Delta Revenue"),
            ("terminal_growth", "Selected Terminal Growth"),
            ("wacc", "Selected WACC"),
        ]:
            ws.cell(row, 1).value = label
            selected_rows[key] = row
            for j in range(2, 7):
                rows = [scenario_rows[(s, key)] for s in ["bear", "base", "bull"]]
                col = get_column_letter(j)
                self.formula(ws, f"{col}{row}", f"=CHOOSE($B$4,{col}{rows[0]},{col}{rows[1]},{col}{rows[2]})")
            row += 1
        row += 1
        self.section(ws, row, "SEGMENT REVENUE", 7)
        self.header_row(ws, row + 1, ["Segment", f"FY{self.hist_years[-1]}A"] + [f"FY{y}E" for y in self.projection_years])
        row += 2
        segment_revenue_rows = []
        for si, segment in enumerate(self.model["segments"]):
            ws.cell(row, 1).value = segment["name"]
            self.input(ws, f"B{row}", segment["revenue"], f"Latest segment revenue input, FY{self.hist_years[-1]}")
            selected_growth_row = selected_rows[f"segment_{si}"]
            for j in range(3, 8):
                prev = f"{get_column_letter(j-1)}{row}"
                growth = f"{get_column_letter(j-1)}{selected_growth_row}"
                self.formula(ws, f"{get_column_letter(j)}{row}", f"={prev}*(1+{growth})")
            segment_revenue_rows.append(row)
            row += 1
        ws.cell(row, 1).value = "Total Revenue"
        for j in range(2, 8):
            col = get_column_letter(j)
            self.formula(ws, f"{col}{row}", f"=SUM({col}{segment_revenue_rows[0]}:{col}{segment_revenue_rows[-1]})")
        self.rows["Revenue Build"] = {"total_revenue": row, **selected_rows}
        self.scenario_rows = scenario_rows

    def build_income_statement(self):
        ws = self.ws("Income Statement")
        self.section(ws, 1, f"Income Statement ({self.model['units']})", 10)
        self.header_row(ws, 3, ["Line item"] + [f"FY{y}A" for y in self.hist_years] + [f"FY{y}E" for y in self.projection_years])
        rows = {"revenue": 4, "growth": 5, "gross_profit": 6, "gross_margin": 7, "ebit": 8, "ebit_margin": 9, "tax": 10, "net_income": 11, "net_margin": 12}
        self.rows["Income Statement"] = rows
        labels = {"revenue": "Revenue", "growth": "Revenue Growth", "gross_profit": "Gross Profit", "gross_margin": "Gross Margin", "ebit": "EBIT", "ebit_margin": "EBIT Margin", "tax": "Tax Expense", "net_income": "Net Income", "net_margin": "Net Margin"}
        for key, r in rows.items():
            ws.cell(r, 1).value = labels[key]
        for i, item in enumerate(self.model["historical"], 2):
            col = get_column_letter(i)
            self.input(ws, f"{col}{rows['revenue']}", item["revenue"], f"{item['source']}, FY{item['year']}")
            if i == 2:
                ws[f"{col}{rows['growth']}"] = "-"
            else:
                self.formula(ws, f"{col}{rows['growth']}", f"={col}{rows['revenue']}/{get_column_letter(i-1)}{rows['revenue']}-1")
            self.input(ws, f"{col}{rows['gross_profit']}", item["gross_profit"], f"{item['source']}, FY{item['year']}")
            self.formula(ws, f"{col}{rows['gross_margin']}", f"={col}{rows['gross_profit']}/{col}{rows['revenue']}")
            self.input(ws, f"{col}{rows['ebit']}", item["ebit"], f"{item['source']}, FY{item['year']}")
            self.formula(ws, f"{col}{rows['ebit_margin']}", f"={col}{rows['ebit']}/{col}{rows['revenue']}")
            tax = max(item["ebit"] - item["net_income"], 0)
            self.input(ws, f"{col}{rows['tax']}", tax, f"{item['source']}, implied FY{item['year']}")
            self.input(ws, f"{col}{rows['net_income']}", item["net_income"], f"{item['source']}, FY{item['year']}")
            self.formula(ws, f"{col}{rows['net_margin']}", f"={col}{rows['net_income']}/{col}{rows['revenue']}")
        for j in range(5, 10):
            col = get_column_letter(j)
            rb_col = get_column_letter(j - 3)
            self.formula(ws, f"{col}{rows['revenue']}", f"='Revenue Build'!{get_column_letter(j-2)}{self.rows['Revenue Build']['total_revenue']}", link=True)
            self.formula(ws, f"{col}{rows['growth']}", f"={col}{rows['revenue']}/{get_column_letter(j-1)}{rows['revenue']}-1")
            self.formula(ws, f"{col}{rows['gross_profit']}", f"={col}{rows['revenue']}*'Revenue Build'!{rb_col}{self.rows['Revenue Build']['gross_margin']}", link=True)
            self.formula(ws, f"{col}{rows['gross_margin']}", f"={col}{rows['gross_profit']}/{col}{rows['revenue']}")
            self.formula(ws, f"{col}{rows['ebit']}", f"={col}{rows['revenue']}*'Revenue Build'!{rb_col}{self.rows['Revenue Build']['ebit_margin']}", link=True)
            self.formula(ws, f"{col}{rows['ebit_margin']}", f"={col}{rows['ebit']}/{col}{rows['revenue']}")
            self.formula(ws, f"{col}{rows['tax']}", f"=MAX(0,{col}{rows['ebit']}*'Revenue Build'!{rb_col}{self.rows['Revenue Build']['tax_rate']})", link=True)
            self.formula(ws, f"{col}{rows['net_income']}", f"={col}{rows['ebit']}-{col}{rows['tax']}")
            self.formula(ws, f"{col}{rows['net_margin']}", f"={col}{rows['net_income']}/{col}{rows['revenue']}")

    def build_balance_sheet(self):
        ws = self.ws("Balance Sheet")
        self.section(ws, 1, f"Balance Sheet ({self.model['units']})", 10)
        self.header_row(ws, 3, ["Line item"] + [f"FY{y}A" for y in self.hist_years] + [f"FY{y}E" for y in self.projection_years])
        rows = {"cash": 4, "nwc": 5, "ppe": 6, "other_assets": 7, "total_assets": 8, "debt": 10, "other_liabilities": 11, "equity": 12, "liab_equity": 13, "balance_check": 14}
        self.rows["Balance Sheet"] = rows
        labels = {"cash": "Cash", "nwc": "Net Working Capital", "ppe": "PP&E", "other_assets": "Other Assets / Plug", "total_assets": "Total Assets", "debt": "Debt", "other_liabilities": "Other Liabilities", "equity": "Equity", "liab_equity": "Total Liabilities + Equity", "balance_check": "Balance Check"}
        for key, r in rows.items():
            ws.cell(r, 1).value = labels[key]
        for i, item in enumerate(self.model["historical"], 2):
            col = get_column_letter(i)
            ppe = item["total_assets"] * 0.35
            other_liabilities = item["total_assets"] - item["debt"] - item["equity"]
            other_assets = item["total_assets"] - item["cash"] - item["nwc"] - ppe
            for key, value in [("cash", item["cash"]), ("nwc", item["nwc"]), ("ppe", ppe), ("other_assets", other_assets), ("debt", item["debt"]), ("other_liabilities", other_liabilities), ("equity", item["equity"])]:
                self.input(ws, f"{col}{rows[key]}", value, f"{item['source']}, FY{item['year']}")
            self.formula(ws, f"{col}{rows['total_assets']}", f"=SUM({col}{rows['cash']}:{col}{rows['other_assets']})")
            self.formula(ws, f"{col}{rows['liab_equity']}", f"=SUM({col}{rows['debt']}:{col}{rows['equity']})")
            self.formula(ws, f"{col}{rows['balance_check']}", f"={col}{rows['total_assets']}-{col}{rows['liab_equity']}")
        for j in range(5, 10):
            col = get_column_letter(j)
            prev = get_column_letter(j - 1)
            rb_col = get_column_letter(j - 3)
            revenue_row = self.rows["Income Statement"]["revenue"]
            net_income_row = self.rows["Income Statement"]["net_income"]
            revenue = f"'Income Statement'!{col}{revenue_row}"
            prev_revenue = f"'Income Statement'!{prev}{revenue_row}"
            da = f"{revenue}*'Revenue Build'!{rb_col}{self.rows['Revenue Build']['da_pct_revenue']}"
            capex = f"{revenue}*'Revenue Build'!{rb_col}{self.rows['Revenue Build']['capex_pct_revenue']}"
            nwc_change = f"({revenue}-{prev_revenue})*'Revenue Build'!{rb_col}{self.rows['Revenue Build']['nwc_pct_delta_revenue']}"
            fcf = f"'Income Statement'!{col}{net_income_row}+{da}-{capex}-{nwc_change}"
            self.formula(ws, f"{col}{rows['cash']}", f"={prev}{rows['cash']}+{fcf}", link=True)
            self.formula(ws, f"{col}{rows['nwc']}", f"={prev}{rows['nwc']}+{nwc_change}", link=True)
            self.formula(ws, f"{col}{rows['ppe']}", f"={prev}{rows['ppe']}+{capex}-{da}", link=True)
            self.formula(ws, f"{col}{rows['debt']}", f"={prev}{rows['debt']}")
            self.formula(ws, f"{col}{rows['other_liabilities']}", f"={prev}{rows['other_liabilities']}")
            self.formula(ws, f"{col}{rows['equity']}", f"={prev}{rows['equity']}+'Income Statement'!{col}{self.rows['Income Statement']['net_income']}", link=True)
            self.formula(ws, f"{col}{rows['liab_equity']}", f"=SUM({col}{rows['debt']}:{col}{rows['equity']})")
            self.formula(ws, f"{col}{rows['other_assets']}", f"={col}{rows['liab_equity']}-{col}{rows['cash']}-{col}{rows['nwc']}-{col}{rows['ppe']}")
            self.formula(ws, f"{col}{rows['total_assets']}", f"=SUM({col}{rows['cash']}:{col}{rows['other_assets']})")
            self.formula(ws, f"{col}{rows['balance_check']}", f"={col}{rows['total_assets']}-{col}{rows['liab_equity']}")

    def build_cash_flow(self):
        ws = self.ws("Cash Flow")
        self.section(ws, 1, f"Cash Flow ({self.model['units']})", 10)
        self.header_row(ws, 3, ["Line item"] + [f"FY{y}A" for y in self.hist_years] + [f"FY{y}E" for y in self.projection_years])
        rows = {"net_income": 4, "da": 5, "capex": 6, "nwc_change": 7, "unlevered_fcf": 8, "beginning_cash": 10, "ending_cash": 11, "cash_tie": 12}
        self.rows["Cash Flow"] = rows
        labels = {"net_income": "Net Income", "da": "(+) D&A", "capex": "(-) CapEx", "nwc_change": "(-) ΔNWC", "unlevered_fcf": "Unlevered FCF", "beginning_cash": "Beginning Cash", "ending_cash": "Ending Cash", "cash_tie": "Cash Tie Check"}
        for key, r in rows.items():
            ws.cell(r, 1).value = labels[key]
        for i, item in enumerate(self.model["historical"], 2):
            col = get_column_letter(i)
            self.formula(ws, f"{col}{rows['net_income']}", f"='Income Statement'!{col}{self.rows['Income Statement']['net_income']}", link=True)
            self.input(ws, f"{col}{rows['da']}", item["da"], f"{item['source']}, FY{item['year']}")
            self.input(ws, f"{col}{rows['capex']}", item["capex"], f"{item['source']}, FY{item['year']}")
            if i == 2:
                self.input(ws, f"{col}{rows['nwc_change']}", 0, f"{item['source']}, opening period")
                self.input(ws, f"{col}{rows['beginning_cash']}", item["cash"], f"{item['source']}, opening cash")
            else:
                self.formula(ws, f"{col}{rows['nwc_change']}", f"='Balance Sheet'!{col}{self.rows['Balance Sheet']['nwc']}-'Balance Sheet'!{get_column_letter(i-1)}{self.rows['Balance Sheet']['nwc']}", link=True)
                self.formula(ws, f"{col}{rows['beginning_cash']}", f"={get_column_letter(i-1)}{rows['ending_cash']}")
            self.formula(ws, f"{col}{rows['unlevered_fcf']}", f"={col}{rows['net_income']}+{col}{rows['da']}-{col}{rows['capex']}-{col}{rows['nwc_change']}")
            self.input(ws, f"{col}{rows['ending_cash']}", item["cash"], f"{item['source']}, reported cash")
            self.formula(ws, f"{col}{rows['cash_tie']}", f"={col}{rows['ending_cash']}-'Balance Sheet'!{col}{self.rows['Balance Sheet']['cash']}", link=True)
        for j in range(5, 10):
            col = get_column_letter(j)
            prev = get_column_letter(j - 1)
            rb_col = get_column_letter(j - 3)
            self.formula(ws, f"{col}{rows['net_income']}", f"='Income Statement'!{col}{self.rows['Income Statement']['net_income']}", link=True)
            self.formula(ws, f"{col}{rows['da']}", f"='Income Statement'!{col}{self.rows['Income Statement']['revenue']}*'Revenue Build'!{rb_col}{self.rows['Revenue Build']['da_pct_revenue']}", link=True)
            self.formula(ws, f"{col}{rows['capex']}", f"='Income Statement'!{col}{self.rows['Income Statement']['revenue']}*'Revenue Build'!{rb_col}{self.rows['Revenue Build']['capex_pct_revenue']}", link=True)
            self.formula(ws, f"{col}{rows['nwc_change']}", f"=('Income Statement'!{col}{self.rows['Income Statement']['revenue']}-'Income Statement'!{prev}{self.rows['Income Statement']['revenue']})*'Revenue Build'!{rb_col}{self.rows['Revenue Build']['nwc_pct_delta_revenue']}", link=True)
            self.formula(ws, f"{col}{rows['unlevered_fcf']}", f"={col}{rows['net_income']}+{col}{rows['da']}-{col}{rows['capex']}-{col}{rows['nwc_change']}")
            self.formula(ws, f"{col}{rows['beginning_cash']}", f"={prev}{rows['ending_cash']}")
            self.formula(ws, f"{col}{rows['ending_cash']}", f"={col}{rows['beginning_cash']}+{col}{rows['unlevered_fcf']}")
            self.formula(ws, f"{col}{rows['cash_tie']}", f"={col}{rows['ending_cash']}-'Balance Sheet'!{col}{self.rows['Balance Sheet']['cash']}", link=True)

    def build_wacc(self):
        ws = self.ws("WACC")
        self.section(ws, 1, "WACC Detail", 5)
        rows = {
            "risk_free": 4, "beta": 5, "erp": 6, "cost_equity": 7,
            "pretax_debt": 10, "tax_rate": 11, "aftertax_debt": 12,
            "price": 15, "shares": 16, "market_cap": 17, "debt": 18, "cash": 19, "net_debt": 20, "ev": 21,
            "equity_weight": 24, "debt_weight": 25, "wacc": 26,
        }
        self.rows["WACC"] = rows
        labels = {
            "risk_free": "Risk-Free Rate", "beta": "Beta", "erp": "Equity Risk Premium", "cost_equity": "Cost of Equity",
            "pretax_debt": "Pre-Tax Cost of Debt", "tax_rate": "Tax Rate", "aftertax_debt": "After-Tax Cost of Debt",
            "price": "Current Share Price", "shares": "Diluted Shares (M)", "market_cap": "Market Cap", "debt": "Debt", "cash": "Cash", "net_debt": "Net Debt", "ev": "Enterprise Value",
            "equity_weight": "Equity Weight", "debt_weight": "Debt Weight", "wacc": "WACC",
        }
        market = self.model["market"]
        for key, label in labels.items():
            ws.cell(rows[key], 1).value = label
        for key in ["risk_free", "beta", "erp", "pretax_debt", "tax_rate", "price", "shares", "debt", "cash"]:
            value = {
                "risk_free": market["risk_free_rate"], "beta": market["beta"], "erp": market["equity_risk_premium"],
                "pretax_debt": market["pretax_cost_of_debt"], "tax_rate": market["tax_rate"], "price": market["price"],
                "shares": market["shares"], "debt": market["debt"], "cash": market["cash"],
            }[key]
            self.input(ws, f"B{rows[key]}", value, "Market/model input; replace with official filings or user-configured market data")
        self.formula(ws, f"B{rows['cost_equity']}", f"=B{rows['risk_free']}+B{rows['beta']}*B{rows['erp']}")
        self.formula(ws, f"B{rows['aftertax_debt']}", f"=B{rows['pretax_debt']}*(1-B{rows['tax_rate']})")
        self.formula(ws, f"B{rows['market_cap']}", f"=B{rows['price']}*B{rows['shares']}")
        self.formula(ws, f"B{rows['net_debt']}", f"=B{rows['debt']}-B{rows['cash']}")
        self.formula(ws, f"B{rows['ev']}", f"=B{rows['market_cap']}+B{rows['net_debt']}")
        self.formula(ws, f"B{rows['equity_weight']}", f"=B{rows['market_cap']}/B{rows['ev']}")
        self.formula(ws, f"B{rows['debt_weight']}", f"=B{rows['net_debt']}/B{rows['ev']}")
        self.formula(ws, f"B{rows['wacc']}", f"=B{rows['cost_equity']}*B{rows['equity_weight']}+B{rows['aftertax_debt']}*B{rows['debt_weight']}")

    def build_dcf(self):
        ws = self.ws("DCF")
        self.section(ws, 1, "DCF Valuation", 10)
        self.header_row(ws, 3, ["Line item"] + [f"FY{y}E" for y in self.projection_years] + ["Terminal"])
        rows = {"fcf": 4, "wacc": 5, "terminal_growth": 6, "period": 7, "discount_factor": 8, "pv_fcf": 9, "terminal_fcf": 11, "terminal_value": 12, "pv_terminal": 13, "pv_fcf_sum": 16, "enterprise_value": 17, "net_debt": 18, "equity_value": 19, "shares": 20, "implied_price": 21, "current_price": 22, "upside": 23}
        self.rows["DCF"] = rows
        labels = {"fcf": "Unlevered FCF", "wacc": "Selected WACC", "terminal_growth": "Terminal Growth", "period": "Discount Period", "discount_factor": "Discount Factor", "pv_fcf": "PV of FCF", "terminal_fcf": "Terminal FCF", "terminal_value": "Terminal Value", "pv_terminal": "PV Terminal Value", "pv_fcf_sum": "Sum of PV FCFs", "enterprise_value": "Enterprise Value", "net_debt": "(-) Net Debt", "equity_value": "Equity Value", "shares": "Shares Outstanding (M)", "implied_price": "IMPLIED PRICE PER SHARE", "current_price": "Current Price", "upside": "Implied Upside/(Downside)"}
        for key, r in rows.items():
            ws.cell(r, 1).value = labels[key]
        for j in range(2, 7):
            col = get_column_letter(j)
            cf_col = get_column_letter(j + 3)
            rb_col = get_column_letter(j)
            self.formula(ws, f"{col}{rows['fcf']}", f"='Cash Flow'!{cf_col}{self.rows['Cash Flow']['unlevered_fcf']}", link=True)
            self.formula(ws, f"{col}{rows['wacc']}", f"='Revenue Build'!{rb_col}{self.rows['Revenue Build']['wacc']}", link=True)
            self.formula(ws, f"{col}{rows['terminal_growth']}", f"='Revenue Build'!{rb_col}{self.rows['Revenue Build']['terminal_growth']}", link=True)
            self.input(ws, f"{col}{rows['period']}", j - 1.5, "Mid-year convention")
            self.formula(ws, f"{col}{rows['discount_factor']}", f"=1/(1+{col}{rows['wacc']})^{col}{rows['period']}")
            self.formula(ws, f"{col}{rows['pv_fcf']}", f"={col}{rows['fcf']}*{col}{rows['discount_factor']}")
        self.formula(ws, f"G{rows['terminal_fcf']}", f"=F{rows['fcf']}*(1+F{rows['terminal_growth']})")
        self.formula(ws, f"G{rows['terminal_value']}", f"=G{rows['terminal_fcf']}/(F{rows['wacc']}-F{rows['terminal_growth']})")
        self.formula(ws, f"G{rows['pv_terminal']}", f"=G{rows['terminal_value']}/(1+F{rows['wacc']})^F{rows['period']}")
        self.formula(ws, f"B{rows['pv_fcf_sum']}", f"=SUM(B{rows['pv_fcf']}:F{rows['pv_fcf']})")
        self.formula(ws, f"B{rows['enterprise_value']}", f"=B{rows['pv_fcf_sum']}+G{rows['pv_terminal']}")
        self.formula(ws, f"B{rows['net_debt']}", f"='WACC'!B{self.rows['WACC']['net_debt']}", link=True)
        self.formula(ws, f"B{rows['equity_value']}", f"=B{rows['enterprise_value']}-B{rows['net_debt']}")
        self.formula(ws, f"B{rows['shares']}", f"='WACC'!B{self.rows['WACC']['shares']}", link=True)
        self.formula(ws, f"B{rows['implied_price']}", f"=B{rows['equity_value']}/B{rows['shares']}")
        self.formula(ws, f"B{rows['current_price']}", f"='WACC'!B{self.rows['WACC']['price']}", link=True)
        self.formula(ws, f"B{rows['upside']}", f"=B{rows['implied_price']}/B{rows['current_price']}-1")

    def build_comps(self):
        ws = self.ws("Comps")
        self.section(ws, 1, "Comparable Companies", 9)
        self.header_row(ws, 3, ["Company", "Ticker", "Market Cap", "EV", "Revenue", "Net Income", "EV/Sales", "P/E", "Source"])
        for r, comp in enumerate(self.model["comps"], 4):
            ws.cell(r, 1).value = comp.get("company")
            ws.cell(r, 2).value = comp.get("ticker")
            for c, key in zip(range(3, 7), ["market_cap", "ev", "revenue", "net_income"]):
                self.input(ws, f"{get_column_letter(c)}{r}", safe_float(comp.get(key)), comp.get("source", "Comps input"))
            self.formula(ws, f"G{r}", f"=IFERROR(D{r}/E{r},0)")
            self.formula(ws, f"H{r}", f"=IFERROR(C{r}/F{r},0)")
            ws.cell(r, 9).value = comp.get("source", "")
        end = 3 + len(self.model["comps"])
        ws.cell(end + 2, 1).value = "Median"
        self.formula(ws, f"G{end+2}", f"=MEDIAN(G4:G{end})")
        self.formula(ws, f"H{end+2}", f"=MEDIAN(H4:H{end})")

    def dcf_sensitivity_formula(self, wacc_ref: str, g_ref: str) -> str:
        r = self.rows["DCF"]
        terms = []
        for idx, col in enumerate(["B", "C", "D", "E", "F"]):
            period = idx + 0.5
            terms.append(f"{col}{r['fcf']}/(1+{wacc_ref})^{period}")
        terminal = f"(F{r['fcf']}*(1+{g_ref})/({wacc_ref}-{g_ref}))/(1+{wacc_ref})^4.5"
        return f"=(SUM({','.join(terms)})+{terminal}-'WACC'!B{self.rows['WACC']['net_debt']})/'WACC'!B{self.rows['WACC']['shares']}"

    def synthetic_dcf_formula(self, growth_ref: str, margin_ref: str, wacc_ref: str | None = None, g_ref: str | None = None) -> str:
        last_revenue_ref = f"'Income Statement'!D{self.rows['Income Statement']['revenue']}"
        tax = f"'Revenue Build'!B{self.rows['Revenue Build']['tax_rate']}"
        da = f"'Revenue Build'!B{self.rows['Revenue Build']['da_pct_revenue']}"
        capex = f"'Revenue Build'!B{self.rows['Revenue Build']['capex_pct_revenue']}"
        nwc = f"'Revenue Build'!B{self.rows['Revenue Build']['nwc_pct_delta_revenue']}"
        wacc = wacc_ref or f"'DCF'!F{self.rows['DCF']['wacc']}"
        tg = g_ref or f"'DCF'!F{self.rows['DCF']['terminal_growth']}"
        fcf_terms = []
        prev_rev = last_revenue_ref
        final_fcf = ""
        for idx in range(5):
            rev = f"({prev_rev}*(1+{growth_ref}))"
            ebit = f"({rev}*{margin_ref})"
            fcf = f"(({ebit})*(1-{tax})+({rev}*{da})-({rev}*{capex})-(({rev}-{prev_rev})*{nwc}))"
            fcf_terms.append(f"{fcf}/(1+{wacc})^{idx+0.5}")
            prev_rev = rev
            final_fcf = fcf
        terminal = f"(({final_fcf})*(1+{tg})/({wacc}-{tg}))/(1+{wacc})^4.5"
        return f"=(SUM({','.join(fcf_terms)})+{terminal}-'WACC'!B{self.rows['WACC']['net_debt']})/'WACC'!B{self.rows['WACC']['shares']}"

    def scenario_dcf_formula(self, scenario_key: str) -> str:
        segment_rows = [
            self.rows["Revenue Build"]["total_revenue"] - len(self.model["segments"]) + idx
            for idx in range(len(self.model["segments"]))
        ]
        segment_values = [f"'Revenue Build'!B{row}" for row in segment_rows]
        fcf_terms = []
        final_fcf = ""
        for idx, col in enumerate(["B", "C", "D", "E", "F"]):
            prev_segment_values = segment_values
            next_segment_values = []
            for si, prev_segment in enumerate(prev_segment_values):
                growth_row = self.scenario_rows[(scenario_key, f"segment_{si}")]
                next_segment_values.append(f"({prev_segment}*(1+'Revenue Build'!{col}{growth_row}))")
            revenue = f"SUM({','.join(next_segment_values)})"
            prev_revenue = f"SUM({','.join(prev_segment_values)})"
            ebit_margin = f"'Revenue Build'!{col}{self.scenario_rows[(scenario_key, 'ebit_margin')]}"
            tax_rate = f"'Revenue Build'!{col}{self.scenario_rows[(scenario_key, 'tax_rate')]}"
            da = f"'Revenue Build'!{col}{self.scenario_rows[(scenario_key, 'da_pct_revenue')]}"
            capex = f"'Revenue Build'!{col}{self.scenario_rows[(scenario_key, 'capex_pct_revenue')]}"
            nwc = f"'Revenue Build'!{col}{self.scenario_rows[(scenario_key, 'nwc_pct_delta_revenue')]}"
            wacc = f"'Revenue Build'!{col}{self.scenario_rows[(scenario_key, 'wacc')]}"
            ebit = f"({revenue}*{ebit_margin})"
            fcf = f"({ebit}*(1-{tax_rate})+({revenue}*{da})-({revenue}*{capex})-(({revenue})-({prev_revenue}))*{nwc})"
            fcf_terms.append(f"{fcf}/(1+{wacc})^{idx+0.5}")
            final_fcf = fcf
            segment_values = next_segment_values
        terminal_wacc = f"'Revenue Build'!F{self.scenario_rows[(scenario_key, 'wacc')]}"
        terminal_growth = f"'Revenue Build'!F{self.scenario_rows[(scenario_key, 'terminal_growth')]}"
        terminal = f"(({final_fcf})*(1+{terminal_growth})/({terminal_wacc}-{terminal_growth}))/(1+{terminal_wacc})^4.5"
        return f"=(SUM({','.join(fcf_terms)})+{terminal}-'WACC'!B{self.rows['WACC']['net_debt']})/'WACC'!B{self.rows['WACC']['shares']}"

    def build_sensitivity(self):
        ws = self.ws("Scenario & Sensitivity")
        self.section(ws, 1, "Bear / Base / Bull Value Range And Sensitivity", 9)
        self.header_row(ws, 3, ["Scenario", "Case Selector", "Implied Price Formula", "Notes"])
        for idx, name in enumerate(["Bear", "Base", "Bull"], 4):
            scenario_key = name.lower()
            ws.cell(idx, 1).value = name
            self.input(ws, f"B{idx}", idx - 3, "Scenario selector")
            self.formula(ws, f"C{idx}", self.scenario_dcf_formula(scenario_key))
            ws.cell(idx, 4).value = "Formula-driven scenario view; DCF tab shows the selected full case."
        start = 8
        self.write_sensitivity_table(ws, start, "WACC vs Terminal Growth", [0.09, 0.095, 0.10, 0.105, 0.11], [0.015, 0.02, 0.025, 0.03, 0.035], lambda r, c: self.dcf_sensitivity_formula(f"$A{r}", f"{c}$9"))
        start += 9
        self.write_sensitivity_table(ws, start, "Revenue Growth vs EBIT Margin", [0.04, 0.07, 0.10, 0.13, 0.16], [0.04, 0.06, 0.08, 0.10, 0.12], lambda r, c: self.synthetic_dcf_formula(f"$A{r}", f"{c}${start+1}"))
        start += 9
        self.write_sensitivity_table(ws, start, "Beta vs Risk-Free Rate", [0.9, 1.05, 1.2, 1.35, 1.5], [0.015, 0.02, 0.025, 0.03, 0.035], lambda r, c: self.dcf_sensitivity_formula(f"(({c}${start+1})+($A{r}*'WACC'!B{self.rows['WACC']['erp']}))", f"'DCF'!F{self.rows['DCF']['terminal_growth']}"))

    def write_sensitivity_table(self, ws, start_row: int, title: str, row_axis: list[float], col_axis: list[float], formula_fn):
        self.section(ws, start_row, title, 7)
        for j, val in enumerate(col_axis, 2):
            self.input(ws, f"{get_column_letter(j)}{start_row+1}", val, f"{title} column axis")
        for i, val in enumerate(row_axis, start_row + 2):
            self.input(ws, f"A{i}", val, f"{title} row axis")
            for j in range(2, 7):
                col = get_column_letter(j)
                self.formula(ws, f"{col}{i}", formula_fn(i, col))
                if i == start_row + 4 and j == 4:
                    ws[f"{col}{i}"].fill = PatternFill("solid", fgColor=MED_BLUE)
                    ws[f"{col}{i}"].font = Font(bold=True, color=FONT_BLACK)

    def build_charts(self):
        ws = self.ws("Charts")
        ws["A1"] = "Charts are linked to model formulas and refresh in Excel/LibreOffice."
        rev_chart = BarChart()
        rev_chart.title = "Revenue by Segment"
        rev_chart.y_axis.title = self.model["units"]
        rb = self.wb["Revenue Build"]
        start_row = self.rows["Revenue Build"]["total_revenue"] - len(self.model["segments"])
        end_row = self.rows["Revenue Build"]["total_revenue"] - 1
        data = Reference(rb, min_col=2, max_col=7, min_row=start_row, max_row=end_row)
        cats = Reference(rb, min_col=2, max_col=7, min_row=start_row - 1, max_row=start_row - 1)
        rev_chart.add_data(data, from_rows=True, titles_from_data=False)
        rev_chart.set_categories(cats)
        ws.add_chart(rev_chart, "A3")
        fcf_chart = LineChart()
        fcf_chart.title = "Unlevered FCF"
        cf = self.wb["Cash Flow"]
        data2 = Reference(cf, min_col=5, max_col=9, min_row=self.rows["Cash Flow"]["unlevered_fcf"], max_row=self.rows["Cash Flow"]["unlevered_fcf"])
        cats2 = Reference(cf, min_col=5, max_col=9, min_row=3, max_row=3)
        fcf_chart.add_data(data2, from_rows=True)
        fcf_chart.set_categories(cats2)
        ws.add_chart(fcf_chart, "A20")

    def build_audit_sheet(self):
        ws = self.ws("Audit")
        self.section(ws, 1, "Model Audit Dashboard", 8)
        self.header_row(ws, 3, ["Check", "Formula / Basis", "Status"])
        checks = [
            ("BS balances", "'Balance Sheet' balance check rows must equal zero", f"=SUM(ABS('Balance Sheet'!E{self.rows['Balance Sheet']['balance_check']}:I{self.rows['Balance Sheet']['balance_check']}))"),
            ("Cash ties out", "'Cash Flow' ending cash equals BS cash", f"=SUM(ABS('Cash Flow'!E{self.rows['Cash Flow']['cash_tie']}:I{self.rows['Cash Flow']['cash_tie']}))"),
            ("WACC valid", "WACC > terminal growth", f"='DCF'!F{self.rows['DCF']['wacc']}-'DCF'!F{self.rows['DCF']['terminal_growth']}"),
            ("DCF output", "Implied price generated", f"='DCF'!B{self.rows['DCF']['implied_price']}"),
        ]
        for r, (name, basis, formula) in enumerate(checks, 4):
            ws.cell(r, 1).value = name
            ws.cell(r, 2).value = basis
            self.formula(ws, f"C{r}", formula)

    def format_workbook(self):
        thin = Side(style="thin", color="B7B7B7")
        for ws in self.wb.worksheets:
            ws.freeze_panes = "B4"
            for row in ws.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                    cell.border = Border(bottom=thin)
                    if isinstance(cell.value, (int, float)):
                        if "Margin" in str(ws.cell(cell.row, 1).value) or "Rate" in str(ws.cell(cell.row, 1).value) or "Growth" in str(ws.cell(cell.row, 1).value) or "WACC" in str(ws.cell(cell.row, 1).value):
                            cell.number_format = "0.0%"
                        else:
                            cell.number_format = '#,##0.0;(#,##0.0);-'
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        if "Margin" in str(ws.cell(cell.row, 1).value) or "Rate" in str(ws.cell(cell.row, 1).value) or "Growth" in str(ws.cell(cell.row, 1).value) or "Upside" in str(ws.cell(cell.row, 1).value):
                            cell.number_format = "0.0%"
                        else:
                            cell.number_format = '#,##0.0;(#,##0.0);-'
            widths = {"A": 34, "B": 16, "C": 16, "D": 16, "E": 16, "F": 16, "G": 16, "H": 16, "I": 16}
            for col, width in widths.items():
                ws.column_dimensions[col].width = width

    def compute_valuation_summary(self) -> dict[str, Any]:
        outputs = {}
        for selector, scenario in [(1, "bear"), (2, "base"), (3, "bull")]:
            outputs[scenario] = self.compute_case_value(scenario)
        values = [outputs[x]["implied_price"] for x in ["bear", "base", "bull"]]
        return {
            "company": self.model["company"],
            "stock_code": self.model["stock_code"],
            "currency": self.model["currency"],
            "units": self.model["units"],
            "cases": outputs,
            "reasonable_value_range": {
                "low": min(values),
                "base": outputs["base"]["implied_price"],
                "high": max(values),
            },
            "methodology": "5-year unlevered DCF with Bear/Base/Bull scenarios, WACC detail, segment revenue build, and sensitivity tables.",
        }

    def compute_case_value(self, scenario: str) -> dict[str, float]:
        assumptions = self.model["scenarios"][scenario]
        segment_revenues = [safe_float(segment.get("revenue")) for segment in self.model["segments"]]
        fcfs = []
        for i in range(5):
            prev_revenue = sum(segment_revenues)
            next_segment_revenues = []
            for si, segment_revenue in enumerate(segment_revenues):
                growth_series = assumptions["segment_growth"][si]
                next_segment_revenues.append(segment_revenue * (1 + growth_series[i]))
            segment_revenues = next_segment_revenues
            revenue = sum(segment_revenues)
            ebit = revenue * assumptions["ebit_margin"][i]
            tax = max(0, ebit * assumptions["tax_rate"][i])
            da = revenue * assumptions["da_pct_revenue"][i]
            capex = revenue * assumptions["capex_pct_revenue"][i]
            nwc = (revenue - prev_revenue) * assumptions["nwc_pct_delta_revenue"][i]
            fcfs.append(ebit - tax + da - capex - nwc)
        wacc = assumptions["wacc"]
        terminal_g = assumptions["terminal_growth"]
        pv_fcf = sum(fcf / ((1 + wacc) ** (i + 0.5)) for i, fcf in enumerate(fcfs))
        terminal = fcfs[-1] * (1 + terminal_g) / max(0.001, wacc - terminal_g)
        pv_terminal = terminal / ((1 + wacc) ** 4.5)
        ev = pv_fcf + pv_terminal
        net_debt = self.model["market"]["debt"] - self.model["market"]["cash"]
        equity = ev - net_debt
        shares = self.model["market"]["shares"] or 1
        raw_price = equity / shares
        price = max(0.0, raw_price)
        return {
            "pv_fcf": round(pv_fcf, 2),
            "pv_terminal": round(pv_terminal, 2),
            "enterprise_value": round(ev, 2),
            "net_debt": round(net_debt, 2),
            "equity_value": round(equity, 2),
            "implied_price": round(price, 2),
            "raw_implied_price": round(raw_price, 2),
            "wacc": round(wacc, 4),
            "terminal_growth": round(terminal_g, 4),
        }

    def run_recalc(self, model_path: Path) -> dict[str, Any]:
        enabled = os.getenv("REPORT_WORKFLOW_RECALC", "").lower() in {"1", "true", "yes", "on"}
        if not enabled:
            return {
                "status": "skipped",
                "reason": "LibreOffice headless recalculation is optional; workbook is configured to recalculate on open.",
            }
        if not RECALC.exists() or not WORKFLOW_PYTHON.exists():
            return {"status": "skipped", "reason": "recalc.py or workflow python not found"}
        try:
            proc = subprocess.run([str(WORKFLOW_PYTHON), str(RECALC), str(model_path), "30"], text=True, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, timeout=45)
            text = proc.stdout.strip()
            try:
                parsed = json.loads(text[text.find("{"):])
                if isinstance(parsed, dict) and "error" in parsed and "status" not in parsed:
                    parsed["status"] = "failed"
                return parsed
            except Exception:
                return {"status": "raw", "returncode": proc.returncode, "output": text[-2000:]}
        except Exception as exc:
            return {"status": "failed", "reason": repr(exc)}

    def audit_markdown(self, model_path: Path, valuation: dict[str, Any], recalc: dict[str, Any]) -> str:
        wb = load_workbook(model_path, data_only=False)
        formulas = 0
        comments = 0
        errors = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formulas += 1
                        if any(err in cell.value for err in ["#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"]):
                            errors.append(f"{ws.title}!{cell.coordinate}: {cell.value}")
                    if cell.comment:
                        comments += 1
        required = ["Cover", "Sources", "Historical Financials", "Revenue Build", "Income Statement", "Balance Sheet", "Cash Flow", "WACC", "DCF", "Comps", "Scenario & Sensitivity", "Charts", "Audit"]
        missing = [x for x in required if x not in wb.sheetnames]
        rng = valuation["reasonable_value_range"]
        lines = [
            "# Financial Model Audit",
            "",
            f"Model: `{model_path.name}`",
            f"Type: Heavy banking equity model - 3-statement + DCF + comps",
            f"Formula cells: {formulas}",
            f"Input comments: {comments}",
            f"Missing required sheets: {', '.join(missing) if missing else 'None'}",
            f"Recalc status: `{recalc.get('status', 'unknown')}`",
            "",
            "## Valuation Range",
            "",
            f"- Bear/Base/Bull reasonable value range: {rng['low']:.2f} / {rng['base']:.2f} / {rng['high']:.2f}",
            "",
            "## Findings",
            "",
        ]
        if missing:
            lines.append("- Critical: required sheet missing.")
        if errors:
            lines.append("- Critical: formula error literals found: " + "; ".join(errors[:10]))
        if recalc.get("status") == "skipped":
            reason = recalc.get("reason", "not requested")
            lines.append(f"- Formula recalculation skipped: {reason}")
        elif recalc.get("status") not in {"success", "raw"}:
            reason = recalc.get("error") or recalc.get("reason") or recalc.get("output") or "no clean success status"
            lines.append(f"- Warning: formula recalc did not return clean success ({reason}); workbook is set to recalculate on open in Excel/LibreOffice.")
        if not missing and not errors:
            lines.append("- No static formula-error literals found. Balance/cash checks are embedded in the Audit sheet.")
        lines.extend([
            "",
            "## Notes",
            "",
            "- Blue font cells are hardcoded inputs; black font cells are formulas; green font cells link across sheets.",
            "- Comps default to placeholders unless `financial_model.comps` is provided or a user-configured data adapter fills them.",
            "- This workbook is a modeling engine. Investment conclusion still requires analyst review.",
        ])
        return "\n".join(lines) + "\n"


def build_financial_model(data: dict[str, Any], sources: list[Any], out_dir: Path) -> dict[str, Any]:
    model = normalize_model_data(data, sources)
    builder = ModelBuilder(model, out_dir)
    return builder.build()


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Build heavy banking financial model for local research workflow")
    parser.add_argument("--input", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    input_data = json.loads(Path(args.input).read_text(encoding="utf-8"))
    result = build_financial_model(input_data, input_data.get("sources", []), Path(args.out))
    print(json.dumps(result, ensure_ascii=False, indent=2))
