#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REQUIRED_FILES = [
    "deep-research-prompt.md",
    "research-pack.json",
    "report.md",
    "financial_model.xlsx",
    "valuation_summary.json",
    "model_audit.md",
]

OPTIONAL_EXPORTS = ["report.docx", "report.pptx", "report.pdf"]

REQUIRED_SHEETS = [
    "Cover",
    "Sources",
    "Historical Financials",
    "Revenue Build",
    "Income Statement",
    "Balance Sheet",
    "Cash Flow",
    "WACC",
    "DCF",
    "Comps",
    "Scenario & Sensitivity",
    "Charts",
    "Audit",
]

EXCEL_ERROR_LITERALS = ["#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NULL!", "#NUM!"]


def validate_run(run_dir: Path, require_exports: bool = False) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []

    if not run_dir.exists():
        return {"status": "failed", "errors": [f"run directory not found: {run_dir}"], "warnings": []}

    for name in REQUIRED_FILES:
        if not (run_dir / name).exists():
            errors.append(f"missing required file: {name}")

    for name in OPTIONAL_EXPORTS:
        exists = (run_dir / name).exists()
        if require_exports and not exists:
            errors.append(f"missing export file: {name}")
        elif not exists:
            warnings.append(f"optional export missing: {name}")

    valuation = {}
    valuation_path = run_dir / "valuation_summary.json"
    if valuation_path.exists():
        valuation = json.loads(valuation_path.read_text(encoding="utf-8"))
        cases = valuation.get("cases", {})
        for case in ["bear", "base", "bull"]:
            if case not in cases:
                errors.append(f"valuation case missing: {case}")
            elif "implied_price" not in cases[case]:
                errors.append(f"valuation case missing implied_price: {case}")
        rng = valuation.get("reasonable_value_range", {})
        if not {"low", "base", "high"}.issubset(rng):
            errors.append("valuation range missing low/base/high")

    model_path = run_dir / "financial_model.xlsx"
    workbook_stats = {}
    if model_path.exists():
        wb = load_workbook(model_path, data_only=False)
        missing_sheets = [name for name in REQUIRED_SHEETS if name not in wb.sheetnames]
        if missing_sheets:
            errors.append("missing workbook sheets: " + ", ".join(missing_sheets))

        formulas = 0
        comments = 0
        formula_errors = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formulas += 1
                        for literal in EXCEL_ERROR_LITERALS:
                            if literal in cell.value:
                                formula_errors.append(f"{ws.title}!{cell.coordinate}:{literal}")
                    if cell.comment:
                        comments += 1

        if formulas < 250:
            warnings.append(f"formula count lower than expected: {formulas}")
        if comments < 100:
            warnings.append(f"input comment count lower than expected: {comments}")
        if formula_errors:
            errors.append("formula error literals found: " + "; ".join(formula_errors[:20]))

        calc = wb.calculation
        if not getattr(calc, "fullCalcOnLoad", False):
            warnings.append("workbook is not marked fullCalcOnLoad")
        if not getattr(calc, "forceFullCalc", False):
            warnings.append("workbook is not marked forceFullCalc")

        workbook_stats = {
            "sheets": len(wb.sheetnames),
            "formulas": formulas,
            "input_comments": comments,
            "fullCalcOnLoad": bool(getattr(calc, "fullCalcOnLoad", False)),
            "forceFullCalc": bool(getattr(calc, "forceFullCalc", False)),
        }

    audit_path = run_dir / "model_audit.md"
    if audit_path.exists():
        audit = audit_path.read_text(encoding="utf-8")
        if "Recalc status: `failed`" in audit:
            errors.append("model audit reports failed recalculation")
        if "Missing required sheets: None" not in audit:
            errors.append("model audit does not confirm required sheets")

    report_path = run_dir / "report.md"
    if report_path.exists():
        report = report_path.read_text(encoding="utf-8")
        phrase_groups = [
            ("valuation section", ["财务与估值框架", "Financial / Valuation Lens"]),
            ("scenario valuation", ["Bear/Base/Bull"]),
            ("sources section", ["引用来源", "Appendix: Sources"]),
        ]
        for label, phrases in phrase_groups:
            if not any(phrase in report for phrase in phrases):
                errors.append(f"report missing {label}")

    return {
        "status": "passed" if not errors else "failed",
        "run_dir": str(run_dir),
        "errors": errors,
        "warnings": warnings,
        "valuation_range": valuation.get("reasonable_value_range", {}),
        "workbook": workbook_stats,
    }


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description="Validate a generated local research workflow run")
    parser.add_argument("run_dir", help="Run directory produced by research_workflow.py")
    parser.add_argument("--require-exports", action="store_true", help="Require DOCX/PPTX/PDF outputs")
    args = parser.parse_args(argv)
    result = validate_run(Path(args.run_dir), args.require_exports)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
