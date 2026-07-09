#!/usr/bin/env python3
from __future__ import annotations

import argparse
import datetime as dt
import html
import io
import json
import os
import re
import shutil
import subprocess
import sys
import textwrap
import urllib.error
import urllib.request
from dataclasses import dataclass, asdict
from pathlib import Path

from financial_model import build_financial_model


ROOT = Path(__file__).resolve().parent
WORKFLOWS_ROOT = ROOT.parent
WORKFLOW_PYTHON = os.getenv("WORKFLOW_PYTHON", sys.executable)
DEFAULT_HUB = Path(os.getenv("AGENT_SKILLS_HUB", str(WORKFLOWS_ROOT.parent))).expanduser()
PROJECT_CONFIG = Path(os.getenv("REPORT_WORKFLOW_PROJECT_CONFIG", str(ROOT / "config" / "project_paths.local.json"))).expanduser()
PROJECT_CONFIG_EXAMPLE = ROOT / "config" / "project_paths.example.json"


@dataclass
class Source:
    id: str
    title: str
    url: str = ""
    publisher: str = ""
    date: str = "unknown"
    note: str = ""
    evidence: str = ""
    status: str = "provided"


def slug(text: str) -> str:
    value = re.sub(r"[^A-Za-z0-9\u4e00-\u9fff]+", "-", text).strip("-")
    return value[:80] or "research-report"


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def load_project_paths() -> dict:
    if PROJECT_CONFIG.exists():
        return read_json(PROJECT_CONFIG)
    if PROJECT_CONFIG_EXAMPLE.exists():
        return read_json(PROJECT_CONFIG_EXAMPLE)
    return {}


def load_input(args: argparse.Namespace) -> dict:
    data = {}
    if args.input:
        data = read_json(Path(args.input))
    if args.company:
        data["company"] = args.company
    if args.industry:
        data["industry"] = args.industry
    if args.topic:
        data["topic"] = args.topic
    if args.template:
        data["template"] = args.template
    company = (data.get("company") or "").strip()
    industry = (data.get("industry") or "").strip()
    topic = (data.get("topic") or "").strip()
    data["company"] = company
    if company and not industry:
        industry = "待识别行业"
    data["industry"] = industry
    if not topic:
        if company:
            topic = f"{company} 基本面、公开财报趋势、产业链位置与估值区间"
        elif industry:
            topic = f"{industry} 行业格局、产业链趋势与估值框架"
        else:
            topic = "未命名主题"
    data["topic"] = topic
    data.setdefault("template", os.getenv("REPORT_WORKFLOW_TEMPLATE", "securities"))
    data.setdefault("language", os.getenv("REPORT_WORKFLOW_LANGUAGE", "zh-CN"))
    data.setdefault("sources", [])
    return data


def extract_pdf_text(raw: bytes, max_pages: int = 8) -> str:
    try:
        from pypdf import PdfReader
    except Exception as exc:
        return f"PDF source retained but pypdf is unavailable: {exc}"
    try:
        reader = PdfReader(io.BytesIO(raw))
        chunks = []
        for page in reader.pages[:max_pages]:
            chunks.append(page.extract_text() or "")
        text = clean_text("\n".join(chunks))
        return text or "PDF source retained; no extractable text found in the first pages."
    except Exception as exc:
        return f"PDF source retained; text extraction failed: {exc}"


def fetch_url(url: str, timeout: int = 20) -> tuple[str, str]:
    req = urllib.request.Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0 research-workflow/0.1",
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        },
    )
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        raw = resp.read(20_000_000)
        ctype = resp.headers.get("content-type", "")
    if "pdf" in ctype.lower() or url.lower().split("?", 1)[0].endswith(".pdf"):
        return extract_pdf_text(raw), "text/plain; extracted-from=pdf"
    charset = "utf-8"
    m = re.search(r"charset=([\w-]+)", ctype)
    if m:
        charset = m.group(1)
    text = raw.decode(charset, "replace")
    return text, ctype


def html_to_text(raw: str) -> tuple[str, str]:
    title_match = re.search(r"<title[^>]*>(.*?)</title>", raw, flags=re.I | re.S)
    title = clean_text(title_match.group(1)) if title_match else ""
    raw = re.sub(r"<script[\s\S]*?</script>", " ", raw, flags=re.I)
    raw = re.sub(r"<style[\s\S]*?</style>", " ", raw, flags=re.I)
    raw = re.sub(r"<[^>]+>", " ", raw)
    return title, clean_text(html.unescape(raw))


def clean_text(text: str) -> str:
    text = re.sub(r"\s+", " ", text or "").strip()
    return text


def normalize_sources(items: list[dict]) -> list[Source]:
    sources: list[Source] = []
    for idx, item in enumerate(items, 1):
        url = item.get("url", "")
        title = item.get("title") or url or f"Source {idx}"
        src = Source(
            id=f"S{idx}",
            title=title,
            url=url,
            publisher=item.get("publisher", ""),
            date=item.get("date", "unknown"),
            note=item.get("note", ""),
            evidence=item.get("evidence") or item.get("text", ""),
            status="provided",
        )
        if url and not src.evidence:
            try:
                raw, ctype = fetch_url(url)
                if "html" in ctype or raw.lstrip().startswith("<"):
                    parsed_title, body = html_to_text(raw)
                    if parsed_title:
                        src.title = parsed_title
                    src.evidence = body[:6000]
                else:
                    src.evidence = clean_text(raw)[:6000]
                src.status = "fetched"
            except Exception as exc:
                src.evidence = f"抓取失败：{exc}"
                src.status = "fetch_failed"
        if len(src.evidence) > 6000:
            src.evidence = src.evidence[:6000] + "..."
        sources.append(src)
    return sources


def research_questions(data: dict) -> list[str]:
    company = data.get("company") or "目标公司"
    industry = data.get("industry") or "目标行业"
    topic = data.get("topic") or company or industry
    return [
        f"{company} 的主营业务、收入结构、利润驱动和关键经营指标是什么？",
        f"{industry} 的市场规模、增长驱动、政策变量和产业链上下游结构是什么？",
        f"{company} 在 {industry} 中的竞争优势、竞争对手和份额变化是什么？",
        f"围绕“{topic}”，有哪些可量化催化剂、估值重估逻辑和反证风险？",
        f"有哪些公司公告、年报、监管披露或行业权威数据可以支撑上述判断？",
    ]


def source_table(sources: list[Source]) -> str:
    if not sources:
        return "暂无外部来源。请先在输入 JSON 的 `sources` 中加入公告、年报、行业报告或网页链接。"
    rows = ["| ID | 标题 | 发布方/日期 | 状态 | 链接 |", "|---|---|---|---|---|"]
    for src in sources:
        pub_date = " / ".join(x for x in [src.publisher, src.date] if x) or "unknown"
        link = f"[link]({src.url})" if src.url else ""
        rows.append(f"| {src.id} | {escape_cell(src.title)} | {escape_cell(pub_date)} | {src.status} | {link} |")
    return "\n".join(rows)


def escape_cell(value: str) -> str:
    return (value or "").replace("|", "\\|").replace("\n", " ")


def evidence_bullets(sources: list[Source], max_items: int = 8) -> str:
    if not sources:
        return "- 待补充资料来源。"
    bullets = []
    for src in sources[:max_items]:
        snippet = src.evidence or src.note or "待提取关键事实。"
        snippet = snippet[:320].strip()
        bullets.append(f"- [{src.id}] {snippet}")
    return "\n".join(bullets)


def fill_report(data: dict, sources: list[Source]) -> str:
    template_name = data.get("template", "securities")
    template_file = ROOT / "templates" / ("consulting_ppt.md" if template_name == "consulting" else "securities_report.md")
    template = template_file.read_text(encoding="utf-8")
    company = data.get("company")
    industry = data.get("industry")
    topic = data.get("topic")
    title_bits = [x for x in [company, industry, topic] if x]
    title = " / ".join(title_bits[:3])
    evidence = evidence_bullets(sources)
    citations = ", ".join(src.id for src in sources[:5]) if sources else "待补充来源"
    fields = {
        "title": title or "本地研报",
        "language": data.get("language", "zh-CN"),
        "investment_view": paragraph(
            "当前版本为资料框架稿。请基于已收集证据判断投资评级、核心逻辑和反证条件。"
            f" 关键证据索引：{citations}。"
        ),
        "executive_summary": evidence,
        "company_overview": paragraph(
            f"围绕 {company or '目标公司'}，优先补充收入结构、利润率、经营指标、客户结构、产能/交付、管理层指引。"
            f" 已抓取来源可作为事实底稿：{citations}。"
        ),
        "industry_chain": paragraph(
            f"围绕 {industry or '目标行业'}，拆分上游资源/设备、中游制造/平台、下游客户/渠道、政策与价格变量。"
        ),
        "competitive_landscape": paragraph(
            "建议按市场份额、产品性能、成本曲线、渠道/客户、资本开支、研发能力和国际化进展做竞品对比。"
        ),
        "valuation_framework": valuation_framework(data),
        "catalysts": "- 业绩超预期或订单兑现。\n- 政策、价格、技术路线或海外扩张变化。\n- 估值锚切换或同业重估。",
        "risks": "- 需求不及预期。\n- 竞争加剧导致价格或毛利率承压。\n- 政策、汇率、供应链、技术路线变化。\n- 当前资料不足导致判断偏差。",
        "strategic_options": "- 增长：寻找高确定性细分场景。\n- 效率：拆解成本、渠道和组织杠杆。\n- 防守：识别替代技术、价格战和客户集中度风险。",
        "action_plan": "- 补齐一手披露资料。\n- 做同业指标表。\n- 建立情景假设。\n- 把关键判断映射到可验证指标。",
        "source_table": source_table(sources),
    }
    return template.format(**fields)


def valuation_framework(data: dict) -> str:
    model_result = data.get("_model_result") or {}
    summary = model_result.get("valuation_summary") or {}
    if not summary:
        return paragraph(
            "券商研报模板下建议使用 PE/PB/EV-EBITDA/DCF/分部估值；咨询 PPT 模板下建议使用场景假设、敏感性和战略价值区间。"
            " 若需要重度投行版，请启用 `--model-depth banking` 生成三表、WACC、DCF、情景估值和模型审计。"
        )
    rng = summary.get("reasonable_value_range", {})
    cases = summary.get("cases", {})
    lines = [
        "本次已生成重度投行版财务模型，包含分业务收入拆分、三张财务报表、WACC 明细、DCF、可比公司表、敏感性分析、图表和模型审计。",
        "",
        f"- Bear/Base/Bull 合理价值区间：{rng.get('low', 'NA')} / {rng.get('base', 'NA')} / {rng.get('high', 'NA')} {summary.get('currency', '')}/股。",
    ]
    for name in ["bear", "base", "bull"]:
        case = cases.get(name)
        if case:
            lines.append(
                f"- {name.title()}：隐含股价 {case.get('implied_price')}，WACC {case.get('wacc'):.1%}，永续增长 {case.get('terminal_growth'):.1%}。"
            )
    lines.extend([
        f"- 方法：{summary.get('methodology', '5 年期 UFCF DCF + 情景分析。')}",
        "- 数据优先级：巨潮公告/年报 PDF、交易所公告/互动；结构化行情和财务字段来自用户已配置的数据源，例如 Wudao MCP、自有 iFinD/同花顺/东方财富或本地 A 股工具；API key 仅从环境变量读取。",
        f"- 模型文件：`{Path(model_result.get('model_path', 'financial_model.xlsx')).name}`；审计文件：`{Path(model_result.get('audit_path', 'model_audit.md')).name}`。",
    ])
    return "\n".join(lines)


def paragraph(text: str) -> str:
    return textwrap.fill(text, width=92)


def write_research_pack(out_dir: Path, data: dict, sources: list[Source]) -> None:
    pack = {
        "input": data,
        "questions": research_questions(data),
        "sources": [asdict(src) for src in sources],
        "project_paths": load_project_paths(),
        "generated_at": dt.datetime.now().isoformat(timespec="seconds"),
    }
    (out_dir / "research-pack.json").write_text(json.dumps(pack, ensure_ascii=False, indent=2), encoding="utf-8")
    prompt_template = (ROOT / "prompts" / "deep_research_prompt.md").read_text(encoding="utf-8")
    prompt = (
        prompt_template
        .replace("{company}", data.get("company") or "")
        .replace("{industry}", data.get("industry") or "")
        .replace("{topic}", data.get("topic") or "")
        .replace("{language}", data.get("language") or "zh-CN")
        .replace("{template}", data.get("template") or "securities")
    )
    (out_dir / "deep-research-prompt.md").write_text(prompt, encoding="utf-8")


def markdown_to_docx(markdown_path: Path, out_path: Path) -> str:
    try:
        from docx import Document
    except Exception as exc:
        return f"skip docx: python-docx unavailable: {exc}"
    doc = Document()
    doc.core_properties.title = markdown_path.stem
    for raw in markdown_path.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        elif line.startswith("> "):
            doc.add_paragraph(line[2:].strip())
        elif line.startswith("|"):
            doc.add_paragraph(line)
        else:
            doc.add_paragraph(line)
    doc.save(out_path)
    return f"wrote {out_path}"


def markdown_to_pptx(markdown_path: Path, out_path: Path, template: str) -> str:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except Exception as exc:
        return f"skip pptx: python-pptx unavailable: {exc}"
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    sections = split_markdown_sections(markdown_path.read_text(encoding="utf-8"))
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = sections[0][0] if sections else markdown_path.stem
    title_slide.placeholders[1].text = "本地研报工作流生成 / editable PPTX"
    for title, body in sections[1:]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title[:80]
        frame = slide.placeholders[1].text_frame
        frame.clear()
        lines = [x.strip("- ").strip() for x in body.splitlines() if x.strip() and not x.startswith("|")]
        for idx, line in enumerate(lines[:8]):
            p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
            p.text = line[:180]
            p.font.size = Pt(18 if template == "consulting" else 15)
            p.level = 0
    prs.save(out_path)
    return f"wrote {out_path}"


def split_markdown_sections(text: str) -> list[tuple[str, str]]:
    sections: list[tuple[str, list[str]]] = []
    current_title = "Report"
    current_body: list[str] = []
    for line in text.splitlines():
        if line.startswith("# "):
            if current_body:
                sections.append((current_title, current_body))
            current_title = line[2:].strip()
            current_body = []
        elif line.startswith("## "):
            if current_body or current_title:
                sections.append((current_title, current_body))
            current_title = line[3:].strip()
            current_body = []
        else:
            current_body.append(line)
    sections.append((current_title, current_body))
    return [(title, "\n".join(body).strip()) for title, body in sections if title]


def markdown_to_html(markdown_path: Path, out_path: Path) -> None:
    css = (ROOT / "templates" / "report_style.css").read_text(encoding="utf-8")
    body = []
    in_ul = False
    for line in markdown_path.read_text(encoding="utf-8").splitlines():
        if line.startswith("# "):
            if in_ul:
                body.append("</ul>")
                in_ul = False
            body.append(f"<h1>{html.escape(line[2:])}</h1>")
        elif line.startswith("## "):
            if in_ul:
                body.append("</ul>")
                in_ul = False
            body.append(f"<h2>{html.escape(line[3:])}</h2>")
        elif line.startswith("- "):
            if not in_ul:
                body.append("<ul>")
                in_ul = True
            body.append(f"<li>{html.escape(line[2:])}</li>")
        elif line.strip():
            if in_ul:
                body.append("</ul>")
                in_ul = False
            body.append(f"<p>{html.escape(line)}</p>")
    if in_ul:
        body.append("</ul>")
    out_path.write_text(f"<!doctype html><html><head><meta charset='utf-8'><style>{css}</style></head><body>{''.join(body)}</body></html>", encoding="utf-8")


def markdown_to_pdf(markdown_path: Path, out_path: Path, html_path: Path) -> str:
    pandoc = shutil.which("pandoc")
    if pandoc:
        proc = subprocess.run([pandoc, str(markdown_path), "-o", str(out_path)], text=True, stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
        if proc.returncode == 0 and out_path.exists():
            return f"wrote {out_path}"
        reason = proc.stdout.strip()[:1200]
    else:
        reason = "pandoc not found"
    markdown_to_html(markdown_path, html_path)
    reportlab_status = markdown_to_reportlab_pdf(markdown_path, out_path)
    if reportlab_status:
        return f"{reportlab_status}; pandoc failed: {reason}"
    return f"pdf fallback: wrote {html_path}; PDF conversion failed: {reason}"


def markdown_to_reportlab_pdf(markdown_path: Path, out_path: Path) -> str:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    except Exception:
        return ""
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        font_name = "STSong-Light"
    except Exception:
        font_name = "Helvetica"
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CJKTitle", parent=styles["Title"], fontName=font_name, fontSize=20, leading=26))
    styles.add(ParagraphStyle(name="CJKHeading", parent=styles["Heading2"], fontName=font_name, fontSize=14, leading=20))
    styles.add(ParagraphStyle(name="CJKBody", parent=styles["BodyText"], fontName=font_name, fontSize=10.5, leading=16))
    story = []
    for raw in markdown_path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            story.append(Spacer(1, 8))
            continue
        if line.startswith("# "):
            story.append(Paragraph(html.escape(line[2:]), styles["CJKTitle"]))
        elif line.startswith("## "):
            story.append(Paragraph(html.escape(line[3:]), styles["CJKHeading"]))
        elif line.startswith("|"):
            story.append(Paragraph(html.escape(line), styles["CJKBody"]))
        elif line.startswith("- "):
            story.append(Paragraph("• " + html.escape(line[2:]), styles["CJKBody"]))
        elif line.startswith("> "):
            story.append(Paragraph(html.escape(line[2:]), styles["CJKBody"]))
        else:
            story.append(Paragraph(html.escape(line), styles["CJKBody"]))
    doc = SimpleDocTemplate(str(out_path), pagesize=A4, rightMargin=42, leftMargin=42, topMargin=48, bottomMargin=48)
    doc.build(story)
    return f"wrote {out_path} via reportlab"


def run(args: argparse.Namespace) -> Path:
    data = load_input(args)
    run_name = slug(data.get("company") or data.get("industry") or data.get("topic") or "report")
    timestamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    out_dir = Path(args.out or ROOT / "runs" / f"{timestamp}-{run_name}")
    out_dir.mkdir(parents=True, exist_ok=True)
    sources = normalize_sources(data.get("sources", []))
    export_log = []
    if args.model_depth != "none":
        model_result = build_financial_model(data, sources, out_dir)
        data["_model_depth"] = args.model_depth
        data["_model_result"] = model_result
        export_log.append(f"wrote {model_result['model_path']}")
        export_log.append(f"wrote {model_result['valuation_path']}")
        export_log.append(f"wrote {model_result['audit_path']}")
    write_research_pack(out_dir, data, sources)
    markdown = fill_report(data, sources)
    md_path = out_dir / "report.md"
    md_path.write_text(markdown, encoding="utf-8")
    exports = set(args.export)
    if "all" in exports:
        exports = {"docx", "pdf", "pptx"}
    if "docx" in exports:
        export_log.append(markdown_to_docx(md_path, out_dir / "report.docx"))
    if "pptx" in exports:
        export_log.append(markdown_to_pptx(md_path, out_dir / "report.pptx", data.get("template", "securities")))
    if "pdf" in exports:
        export_log.append(markdown_to_pdf(md_path, out_dir / "report.pdf", out_dir / "report.html"))
    (out_dir / "export-log.txt").write_text("\n".join(export_log) + "\n", encoding="utf-8")
    print(out_dir)
    return out_dir


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Local equity/industry research report workflow")
    parser.add_argument("--input", help="JSON input file with company/industry/topic/sources")
    parser.add_argument("--company", help="Company name")
    parser.add_argument("--industry", help="Industry theme")
    parser.add_argument("--topic", help="Research topic")
    parser.add_argument("--template", choices=["securities", "consulting"], help="Output template")
    parser.add_argument("--out", help="Output run directory")
    parser.add_argument("--export", nargs="+", default=["all"], choices=["all", "docx", "pdf", "pptx", "none"], help="Export formats")
    parser.add_argument("--model-depth", choices=["none", "standard", "banking"], default=os.getenv("REPORT_WORKFLOW_MODEL_DEPTH", "banking"), help="Financial model depth; banking builds the full 3-statement/WACC/DCF/audit model")
    return parser.parse_args(argv)


if __name__ == "__main__":
    ns = parse_args(sys.argv[1:])
    if "none" in ns.export:
        ns.export = []
    run(ns)
