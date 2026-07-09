#!/usr/bin/env python3
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import shlex
import textwrap
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
WORKFLOWS_ROOT = ROOT.parent
HUB = Path(os.getenv("AGENT_SKILLS_HUB", str(WORKFLOWS_ROOT.parent))).expanduser()
OUTPUT_WORKFLOW = Path(
    os.getenv("LOCAL_RESEARCH_WORKFLOW", str(WORKFLOWS_ROOT / "local-research-workflow" / "research_workflow.py"))
).expanduser()
WORKFLOW_PYTHON = os.getenv("WORKFLOW_PYTHON", "python")


EVIDENCE_LEVELS = {
    "A": "公告/年报/季报/交易所互动/SEC/IR/业绩会明确确认",
    "B": "券商研报、行业报告、公开数据多源交叉验证",
    "C": "新闻、市场传闻、概念标签、单一非官方来源",
    "D": "公司否认、证据不足、长期无兑现或已被证伪",
}

FACTOR_ROWS = [
    ("product_hardness", "产品硬度", 12, "是否卡在高价值量/高壁垒环节，而不是软概念。"),
    ("bottleneck", "瓶颈强度", 15, "供给是否紧缺、扩产是否慢、良率/认证是否构成约束。"),
    ("recognizability", "辨识度", 14, "市场能否一句话把公司贴成主线核心标的。"),
    ("leader_ladder", "龙头/梯队", 12, "板块龙一/龙二/补涨/跟风的位置。"),
    ("sentiment_funds", "情绪资金", 10, "涨停、热榜、主力资金、融资/北向、龙虎榜等。"),
    ("price_volume_position", "位置量价", 10, "20/60/120 日涨幅、成交额分位、高位天量或深调低吸。"),
    ("catalyst", "催化剂", 8, "业绩、订单、产能投产、政策、大会、新品等接力点。"),
    ("disproof_difficulty", "证伪难度", 9, "订单/客户/技术路线是否容易被证伪。"),
    ("chip_risk", "筹码风险", 10, "解禁、减持、质押、流通市值、持有人结构。"),
]


@dataclass
class Source:
    id: str
    title: str
    url: str = ""
    publisher: str = ""
    date: str = "unknown"
    note: str = ""
    evidence: str = ""
    evidence_level: str = "C"
    claim: str = ""


def slug(text: str) -> str:
    value = re.sub(r"[^A-Za-z0-9\u4e00-\u9fff]+", "-", text or "").strip("-")
    return value[:80] or "analysis"


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def shell_join(parts: list[str | Path]) -> str:
    return " ".join(shlex.quote(str(part)) for part in parts)


def portable_path(path: Path) -> str:
    resolved = path.resolve()
    try:
        return resolved.relative_to(WORKFLOWS_ROOT.resolve()).as_posix()
    except ValueError:
        return str(resolved)


def handoff_command(out_dir: Path) -> str:
    return shell_join(
        [
            WORKFLOW_PYTHON,
            portable_path(OUTPUT_WORKFLOW),
            "--input",
            portable_path(out_dir / "export_input.json"),
            "--export",
            "all",
            "--model-depth",
            "banking",
        ]
    )


def load_input(args: argparse.Namespace) -> dict[str, Any]:
    data: dict[str, Any] = {}
    if args.input:
        data = read_json(Path(args.input))
    for key, value in {
        "company": args.company,
        "stock_code": args.stock_code,
        "industry": args.industry,
        "topic": args.topic,
        "market": args.market,
    }.items():
        if value:
            data[key] = value
    company = (data.get("company") or "").strip()
    industry = (data.get("industry") or "").strip()
    topic = (data.get("topic") or "").strip()
    data["company"] = company
    data["market"] = data.get("market") or "A-share"
    data["language"] = data.get("language") or "zh-CN"
    data["industry"] = industry or "待识别行业"
    data["topic"] = topic or f"{company or data['industry']} 分析版工作流：基本面、产业链、竞品、替代风险、A股主题因子与估值"
    data.setdefault("sources", [])
    data.setdefault("claims", [])
    data.setdefault("competitors", [])
    data.setdefault("substitution_risks", [])
    data.setdefault("catalysts", [])
    data.setdefault("tracking_indicators", [])
    data.setdefault("factor_scores", {})
    data.setdefault("financial_model", {})
    return data


def normalize_sources(items: list[dict[str, Any]]) -> list[Source]:
    sources: list[Source] = []
    for idx, item in enumerate(items, 1):
        level = (item.get("evidence_level") or item.get("level") or "C").upper()
        if level not in EVIDENCE_LEVELS:
            level = "C"
        sources.append(
            Source(
                id=item.get("id") or f"S{idx}",
                title=item.get("title") or item.get("url") or f"Source {idx}",
                url=item.get("url", ""),
                publisher=item.get("publisher", ""),
                date=item.get("date", "unknown"),
                note=item.get("note", ""),
                evidence=item.get("evidence") or item.get("text") or "",
                evidence_level=level,
                claim=item.get("claim", ""),
            )
        )
    return sources


def build_claims(data: dict[str, Any], sources: list[Source]) -> list[dict[str, Any]]:
    claims = list(data.get("claims") or [])
    if claims:
        for idx, claim in enumerate(claims, 1):
            claim.setdefault("id", f"C{idx}")
            claim.setdefault("evidence_level", "C")
            claim.setdefault("status", "待核验")
        return claims
    generated = []
    for idx, src in enumerate(sources, 1):
        generated.append(
            {
                "id": f"C{idx}",
                "claim": src.claim or src.note or f"来自 {src.title} 的待提取事实",
                "evidence_level": src.evidence_level,
                "source_ids": [src.id],
                "status": "待核验" if src.evidence_level in {"C", "D"} else "可引用",
                "analyst_note": src.evidence[:240] if src.evidence else "",
            }
        )
    return generated


def default_factor_scores(data: dict[str, Any]) -> dict[str, Any]:
    scores = dict(data.get("factor_scores") or {})
    for key, _, _, _ in FACTOR_ROWS:
        item = scores.get(key)
        if isinstance(item, (int, float)):
            scores[key] = {"score": item, "evidence_level": "C", "note": ""}
        elif not isinstance(item, dict):
            scores[key] = {"score": "", "evidence_level": "C", "note": "待分析师或 agent 填入"}
        else:
            item.setdefault("score", "")
            item.setdefault("evidence_level", "C")
            item.setdefault("note", "")
    return scores


def md_table(headers: list[str], rows: list[list[Any]]) -> str:
    out = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    for row in rows:
        out.append("| " + " | ".join(escape_cell(str(x)) for x in row) + " |")
    return "\n".join(out)


def escape_cell(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", " ")


def evidence_level_table() -> str:
    return md_table(["等级", "定义"], [[k, v] for k, v in EVIDENCE_LEVELS.items()])


def write_analysis_markdown(out_dir: Path, data: dict[str, Any], sources: list[Source], claims: list[dict[str, Any]], factor_scores: dict[str, Any]) -> None:
    company = data.get("company") or "目标公司"
    stock_code = data.get("stock_code") or "待识别"
    industry = data.get("industry") or "待识别行业"
    claim_rows = [
        [
            c.get("id", ""),
            c.get("claim", ""),
            c.get("evidence_level", "C"),
            ",".join(c.get("source_ids", [])),
            c.get("status", "待核验"),
        ]
        for c in claims[:20]
    ]
    factor_rows = []
    weighted_total_formula_hint = []
    for key, name, weight, definition in FACTOR_ROWS:
        item = factor_scores.get(key, {})
        factor_rows.append([name, weight, item.get("score", ""), item.get("evidence_level", "C"), item.get("note", definition)])
        weighted_total_formula_hint.append(f"{name}×{weight}%")
    text = f"""# {company} 分析版工作流

生成时间：{dt.datetime.now().isoformat(timespec="seconds")}

## 0. 使用定位

这是上游分析包，不负责设计输出。它的任务是把公司研究、A股主题因子、竞品格局、替代风险、财务估值和跟踪计划沉淀为结构化输入，后续可直接交给 `local-research-workflow` 生成券商研报、DOCX、PDF、PPTX 或设计版输出。

## 1. 公司识别

- 公司：{company}
- 代码：{stock_code}
- 市场：{data.get("market", "A-share")}
- 行业：{industry}
- 主题：{data.get("topic")}

## 2. 推荐 Skill 编排

- `agent-skill-stack-router`：总路由。
- `local-analysis-workflow`：本分析包入口。
- `investment-sentiment`：A股短线情绪、题材强度、涨停梯队、资金和位置。
- `finance-stack`：财务模型、DCF、可比公司、Bear/Base/Bull。
- `business-framework-stack`：竞品、行业格局、替代风险、Porter/价值链。
- `local-research-workflow`：下游输出版工作流。

## 3. 证据等级

{evidence_level_table()}

## 4. 证据台账摘要

{md_table(["ID", "核心判断", "等级", "来源", "状态"], claim_rows) if claim_rows else "暂无证据台账。请先补充公告、年报、交易所互动、研报或行业资料。"}

## 5. 公司基本面分析框架

- 收入结构：按业务/产品/区域/客户拆分。
- 利润质量：毛利率、费用率、经营杠杆、减值、一次性损益。
- 现金流：经营现金流、资本开支、营运资本、应收和存货。
- 资产负债：有息负债、商誉、质押、担保、流动性。
- 管理层指引：订单、产能、价格、客户、海外化。

## 6. 产业链位置

- 上游：原材料、设备、核心零部件、关键供应商。
- 中游：制造/平台/系统集成/服务。
- 下游：客户、应用场景、需求周期、议价权。
- 价值量：单机/单柜/单项目价值量和毛利率位置。
- 瓶颈：扩产周期、良率、认证、客户绑定、国产替代稀缺性。

## 7. A股主题因子

综合分建议按：{", ".join(weighted_total_formula_hint)}。估值在主题爆发期降级为风控变量，位置量价和筹码风险承担择时职责。

{md_table(["因子", "权重", "评分(0-5)", "证据等级", "说明"], factor_rows)}

## 8. 竞品分析

详见 `competitor_map.md`。重点区分直接竞品、间接竞品、海外竞品、客户自研和上下游整合。

## 9. 未来行业格局

详见 `industry_structure.md`。重点回答集中度、供需、价格战、技术路线、利润池迁移和中国公司在全球链条里的位置。

## 10. 替代风险

详见 `substitution_risk.md`。必须给出替代来源、替代强度、影响时间点和公司防守能力。

## 11. 财务模型与估值

- 历史数据优先级：巨潮公告/年报 PDF/交易所公告/交易所互动；结构化行情和财务字段来自用户已配置的数据源，例如 Wudao MCP、自有 iFinD/同花顺/东方财富或本地 A 股工具。
- 输出工作流接收字段：`financial_model.historical`、`segments`、`market`、`comps`、`scenarios`。
- 下游 `local-research-workflow` 会生成 `financial_model.xlsx`、`valuation_summary.json` 和 `model_audit.md`。

## 12. 反方研究

- 核心逻辑哪里可能错？
- 哪个数据会证伪？
- 竞品是否更强？
- 新产品或新技术是否会绕开当前产品？
- 市场是否已经过度定价？
- 如果情绪退潮，先杀估值、位置还是业绩？

## 13. 催化剂与跟踪

详见 `catalyst_calendar.md` 和 `tracking_plan.md`。

## 14. 下游交付

本次已生成 `export_input.json`，后续可直接运行：

```bash
{handoff_command(out_dir)}
```
"""
    (out_dir / "analysis.md").write_text(text, encoding="utf-8")


def write_competitor_map(out_dir: Path, data: dict[str, Any]) -> None:
    competitors = data.get("competitors") or []
    rows = []
    for item in competitors:
        rows.append([
            item.get("name", ""),
            item.get("ticker", ""),
            item.get("type", "直接竞品"),
            item.get("product_overlap", ""),
            item.get("customer_overlap", ""),
            item.get("advantage", ""),
            item.get("risk_to_target", ""),
            item.get("evidence_level", "C"),
        ])
    if not rows:
        rows = [
            ["待补充", "", "直接竞品", "同产品/同客户/同环节", "待核验", "待核验", "抢份额/压价格/抢认证", "C"],
            ["待补充", "", "间接竞品", "替代技术/替代材料/替代商业模式", "待核验", "待核验", "绕开当前利润池", "C"],
            ["待补充", "", "海外竞品", "全球供应链可比公司", "待核验", "待核验", "客户优先级和认证壁垒", "C"],
        ]
    text = f"""# 竞品分析

## 竞品矩阵

{md_table(["名称", "代码", "类型", "产品重叠", "客户重叠", "优势", "对目标公司的威胁", "证据等级"], rows)}

## 分析问题

- 谁是同产品、同客户、同环节的直接竞品？
- 谁是替代技术、替代材料、替代架构的间接竞品？
- 海外竞品是否更接近产业链真实定价？
- 客户是否可能自研或扶持第二供应商？
- 上游是否可能前向整合，下游是否可能压价？
- 未来利润池在产品、材料、设备、系统还是服务层？
"""
    (out_dir / "competitor_map.md").write_text(text, encoding="utf-8")


def write_industry_structure(out_dir: Path, data: dict[str, Any]) -> None:
    text = f"""# 未来行业格局

## 核心问题

- 行业集中度未来提高还是分散？
- 龙头份额是否继续提升，还是二线厂商补涨？
- 供给扩张时间表是什么？瓶颈会在 1 年、1-3 年还是 3-5 年缓解？
- 下游需求是周期性拉动、政策拉动，还是技术迭代带来的持续替换？
- 价格战概率多高？毛利率会维持、上行还是下行？
- 中国公司在全球链条里的位置是替代、配套、跟随还是主导？

## Porter 五力检查

- 现有竞争：份额、价格、毛利率、产能利用率。
- 新进入者：认证周期、资本开支、良率、客户验证。
- 供应商议价：关键材料/设备/知识产权是否受限。
- 客户议价：客户集中度、双供策略、自研能力。
- 替代品：新技术路线、新材料、新架构、新商业模式。

## 行业阶段判断

```text
导入期：技术路线未定，证据等级低，弹性来自想象。
成长期：需求兑现，瓶颈涨价，龙头和核心供应商占优。
扩产期：供给增加，价格和毛利率开始分化。
成熟期：估值回到现金流和份额，主题溢价收敛。
替代期：新产品/新架构迁移利润池，原主线降级。
```
"""
    (out_dir / "industry_structure.md").write_text(text, encoding="utf-8")


def write_substitution_risk(out_dir: Path, data: dict[str, Any]) -> None:
    risks = data.get("substitution_risks") or []
    rows = []
    for item in risks:
        rows.append([
            item.get("risk", ""),
            item.get("source", ""),
            item.get("timeframe", ""),
            item.get("intensity", ""),
            item.get("impact", ""),
            item.get("defense", ""),
            item.get("evidence_level", "C"),
        ])
    if not rows:
        rows = [
            ["技术路线替代", "新架构/新材料/新工艺", "1-3年", "局部替代", "利润池迁移", "客户认证、第二曲线、成本优势", "C"],
            ["客户自研", "核心客户内部化", "3-5年", "边际侵蚀", "议价权下降", "多客户和服务能力", "C"],
            ["供给扩产", "行业产能集中释放", "1-3年", "价格下行", "毛利率承压", "高端产品和良率壁垒", "C"],
        ]
    text = f"""# 替代风险与时间表

{md_table(["风险", "来源", "时间点", "强度", "影响", "防守能力", "证据等级"], rows)}

## 判断规则

- 1 年内：直接影响交易节奏和估值弹性。
- 1-3 年：影响中期利润率、产能投放和估值中枢。
- 3-5 年：影响长期终局，但短期可能仍有景气兑现窗口。
- 5 年以上：作为战略风险，不轻易否定当前主线。

## 必答问题

- 替代是边际侵蚀、局部替代，还是全面替代？
- 替代前是否仍有涨价、满产、国产替代或份额提升窗口？
- 公司是否有第二曲线，还是只暴露在旧产品利润池？
"""
    (out_dir / "substitution_risk.md").write_text(text, encoding="utf-8")


def write_catalyst_calendar(out_dir: Path, data: dict[str, Any]) -> None:
    catalysts = data.get("catalysts") or []
    rows = []
    for item in catalysts:
        rows.append([
            item.get("date", "待定"),
            item.get("event", ""),
            item.get("type", ""),
            item.get("expected_impact", ""),
            item.get("source_ids", ""),
            item.get("status", "待跟踪"),
        ])
    if not rows:
        rows = [
            ["待定", "年报/季报/业绩预告", "业绩", "验证收入、利润率和现金流", "", "待跟踪"],
            ["待定", "订单/定点/客户认证", "经营", "验证产品硬度和瓶颈兑现", "", "待跟踪"],
            ["待定", "产能投放/扩产", "供给", "判断瓶颈缓解时间", "", "待跟踪"],
            ["待定", "政策/大会/新品", "主题", "判断情绪接力和估值重估", "", "待跟踪"],
            ["待定", "解禁/减持/质押", "风险", "判断筹码风险", "", "待跟踪"],
        ]
    (out_dir / "catalyst_calendar.md").write_text("# 催化剂日历\n\n" + md_table(["日期", "事件", "类型", "影响", "来源", "状态"], rows) + "\n", encoding="utf-8")


def write_tracking_plan(out_dir: Path, data: dict[str, Any]) -> None:
    indicators = data.get("tracking_indicators") or []
    rows = []
    for item in indicators:
        rows.append([
            item.get("indicator", ""),
            item.get("why_it_matters", ""),
            item.get("source", ""),
            item.get("frequency", ""),
            item.get("trigger", ""),
        ])
    if not rows:
        rows = [
            ["收入增速/分业务收入", "验证主业兑现", "年报/季报/用户自配数据源", "季度", "连续两个季度低于假设"],
            ["毛利率/价格", "验证瓶颈和定价权", "财报/调研/产业数据", "季度/月度", "毛利率拐头下行"],
            ["订单/客户认证", "验证产品硬度", "公告/互动/研报", "事件驱动", "客户或订单不及预期"],
            ["竞品扩产/新品", "验证替代风险", "公告/行业新闻", "月度", "竞品进入核心客户"],
            ["位置量价/资金", "验证A股交易节奏", "Wudao MCP 或本地 A 股工具", "日/周", "高位天量滞涨或放量跌破"],
        ]
    text = "# 跟踪计划\n\n" + md_table(["指标", "重要性", "来源", "频率", "触发条件"], rows) + "\n"
    (out_dir / "tracking_plan.md").write_text(text, encoding="utf-8")


def write_deep_analysis_prompt(out_dir: Path, data: dict[str, Any]) -> None:
    prompt = f"""# Deep Analysis Prompt

请为 `{data.get("company") or data.get("industry")}` 做分析版工作流，不产出设计文件。

## 数据优先级

1. A股优先：巨潮公告、年报 PDF、交易所公告、交易所互动；行情、资金和结构化字段来自用户已配置的数据源，例如 Wudao MCP、自有 iFinD/同花顺/东方财富或本地 A 股工具。
2. 美股/海外补充：SEC、IR、10-K/10-Q、业绩会、公开财报趋势。
3. 行业补充：SemiAnalysis、行业文章、咨询资料、产业链公开数据。
4. 不用软文覆盖官方披露；传闻只能标为 C 级证据。

## 必须输出

- evidence_ledger：每个关键判断都标 A/B/C/D 证据等级。
- competitor_map：直接竞品、间接竞品、海外竞品、客户自研、上下游整合。
- industry_structure：未来行业格局、供需、价格战、利润池迁移。
- substitution_risk：替代来源、时间点、强度、影响、防守能力。
- factor_scores：A股主题因子，0-5 分，含产品、瓶颈、辨识度、龙头、情绪、位置、催化、证伪、筹码。
- financial_model：历史数据、分业务、市场、comps、scenarios，给下游 `local-research-workflow` 使用。

## 推荐本地 skill / tool

- `investment-sentiment`：short_term_emotion、concept_ranking、concept_stocks、hot_sectors、capital_flow、kline。
- `finance-stack`：DCF、三表、comps、Bear/Base/Bull。
- `business-framework-stack`：竞品、Porter、价值链、替代风险。
- `local-research-workflow`：下游输出。
"""
    (out_dir / "deep-analysis-prompt.md").write_text(prompt, encoding="utf-8")


def write_factor_workbook(out_dir: Path, data: dict[str, Any], sources: list[Source], claims: list[dict[str, Any]], factor_scores: dict[str, Any]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Factor Score"
    headers = ["因子", "权重", "评分(0-5)", "加权分", "证据等级", "说明"]
    ws.append(headers)
    for idx, (key, name, weight, definition) in enumerate(FACTOR_ROWS, start=2):
        item = factor_scores.get(key, {})
        score = item.get("score", "")
        ws.append([name, weight, score, f'=IF(ISNUMBER(C{idx}),B{idx}*C{idx}/5,"")', item.get("evidence_level", "C"), item.get("note") or definition])
    total_row = len(FACTOR_ROWS) + 2
    ws.append(["总分", 100, "", f"=SUM(D2:D{total_row-1})", "", "70+ 主线核心；55-70 观察/波段；<55 降级或仅跟踪"])
    style_sheet(ws)

    ev = wb.create_sheet("Evidence Ledger")
    ev.append(["ID", "判断", "证据等级", "来源", "状态", "备注"])
    for item in claims:
        ev.append([item.get("id"), item.get("claim"), item.get("evidence_level"), ",".join(item.get("source_ids", [])), item.get("status"), item.get("analyst_note", "")])
    style_sheet(ev)

    src_ws = wb.create_sheet("Sources")
    src_ws.append(["ID", "标题", "发布方", "日期", "等级", "URL", "备注"])
    for src in sources:
        src_ws.append([src.id, src.title, src.publisher, src.date, src.evidence_level, src.url, src.note])
    style_sheet(src_ws)

    comp = wb.create_sheet("Competitors")
    comp.append(["名称", "代码", "类型", "产品重叠", "客户重叠", "优势", "威胁", "证据等级"])
    for item in data.get("competitors") or []:
        comp.append([item.get("name"), item.get("ticker"), item.get("type"), item.get("product_overlap"), item.get("customer_overlap"), item.get("advantage"), item.get("risk_to_target"), item.get("evidence_level", "C")])
    style_sheet(comp)

    sub = wb.create_sheet("Substitution Risk")
    sub.append(["风险", "来源", "时间点", "强度", "影响", "防守能力", "证据等级"])
    for item in data.get("substitution_risks") or []:
        sub.append([item.get("risk"), item.get("source"), item.get("timeframe"), item.get("intensity"), item.get("impact"), item.get("defense"), item.get("evidence_level", "C")])
    style_sheet(sub)

    cat = wb.create_sheet("Catalysts")
    cat.append(["日期", "事件", "类型", "影响", "来源", "状态"])
    for item in data.get("catalysts") or []:
        cat.append([item.get("date"), item.get("event"), item.get("type"), item.get("expected_impact"), item.get("source_ids"), item.get("status")])
    style_sheet(cat)

    track = wb.create_sheet("Tracking")
    track.append(["指标", "重要性", "来源", "频率", "触发条件"])
    for item in data.get("tracking_indicators") or []:
        track.append([item.get("indicator"), item.get("why_it_matters"), item.get("source"), item.get("frequency"), item.get("trigger")])
    style_sheet(track)

    handoff = wb.create_sheet("Output Handoff")
    handoff.append(["字段", "值"])
    handoff.append(["export_input", str(out_dir / "export_input.json")])
    handoff.append(["output_workflow", portable_path(OUTPUT_WORKFLOW)])
    handoff.append(["command", handoff_command(out_dir)])
    style_sheet(handoff)

    wb.save(out_dir / "factor_score.xlsx")


def style_sheet(ws) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = min(42, max(12, len(str(ws.cell(1, col).value or "")) + 8))
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_export_input(data: dict[str, Any], sources: list[Source], claims: list[dict[str, Any]], factor_scores: dict[str, Any]) -> dict[str, Any]:
    clean_sources = []
    for src in sources:
        item = asdict(src)
        item["note"] = f"[Evidence {src.evidence_level}] {src.note}".strip()
        clean_sources.append(item)
    return {
        "company": data.get("company"),
        "stock_code": data.get("stock_code", ""),
        "industry": data.get("industry"),
        "topic": data.get("topic"),
        "template": data.get("template", "securities"),
        "language": data.get("language", "zh-CN"),
        "sources": clean_sources,
        "financial_model": data.get("financial_model", {}),
        "analysis_context": {
            "workflow": "local-analysis-workflow",
            "market": data.get("market", "A-share"),
            "evidence_levels": EVIDENCE_LEVELS,
            "claims": claims,
            "factor_scores": factor_scores,
            "competitors": data.get("competitors", []),
            "substitution_risks": data.get("substitution_risks", []),
            "catalysts": data.get("catalysts", []),
            "tracking_indicators": data.get("tracking_indicators", []),
        },
    }


def write_analysis_pack(out_dir: Path, data: dict[str, Any], sources: list[Source], claims: list[dict[str, Any]], factor_scores: dict[str, Any], export_input: dict[str, Any]) -> None:
    pack = {
        "generated_at": dt.datetime.now().isoformat(timespec="seconds"),
        "input": data,
        "sources": [asdict(src) for src in sources],
        "claims": claims,
        "factor_scores": factor_scores,
        "output_files": {
            "analysis": "analysis.md",
            "evidence_ledger": "evidence_ledger.json",
            "competitor_map": "competitor_map.md",
            "industry_structure": "industry_structure.md",
            "substitution_risk": "substitution_risk.md",
            "factor_score": "factor_score.xlsx",
            "catalyst_calendar": "catalyst_calendar.md",
            "tracking_plan": "tracking_plan.md",
            "export_input": "export_input.json",
        },
        "handoff": {
            "output_workflow": portable_path(OUTPUT_WORKFLOW),
            "command": handoff_command(out_dir),
        },
        "export_input_preview": export_input,
    }
    (out_dir / "analysis-pack.json").write_text(json.dumps(pack, ensure_ascii=False, indent=2), encoding="utf-8")


def run(args: argparse.Namespace) -> Path:
    data = load_input(args)
    sources = normalize_sources(data.get("sources", []))
    claims = build_claims(data, sources)
    factor_scores = default_factor_scores(data)
    run_name = slug(data.get("company") or data.get("industry") or data.get("topic"))
    timestamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    out_dir = Path(args.out) if args.out else ROOT / "runs" / f"{timestamp}-{run_name}-analysis"
    out_dir.mkdir(parents=True, exist_ok=True)

    write_deep_analysis_prompt(out_dir, data)
    write_analysis_markdown(out_dir, data, sources, claims, factor_scores)
    write_competitor_map(out_dir, data)
    write_industry_structure(out_dir, data)
    write_substitution_risk(out_dir, data)
    write_catalyst_calendar(out_dir, data)
    write_tracking_plan(out_dir, data)
    write_factor_workbook(out_dir, data, sources, claims, factor_scores)
    export_input = build_export_input(data, sources, claims, factor_scores)
    (out_dir / "export_input.json").write_text(json.dumps(export_input, ensure_ascii=False, indent=2), encoding="utf-8")
    (out_dir / "evidence_ledger.json").write_text(json.dumps({"levels": EVIDENCE_LEVELS, "claims": claims, "sources": [asdict(src) for src in sources]}, ensure_ascii=False, indent=2), encoding="utf-8")
    (out_dir / "handoff_command.txt").write_text(handoff_command(out_dir) + "\n", encoding="utf-8")
    write_analysis_pack(out_dir, data, sources, claims, factor_scores, export_input)
    print(out_dir)
    return out_dir


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Local analysis-first workflow for A-share company/industry research")
    parser.add_argument("--input", help="JSON input file with company, sources, factors, competitors and financial_model")
    parser.add_argument("--company", help="Company name, e.g. 埃斯顿")
    parser.add_argument("--stock-code", help="Stock code, e.g. 002747.SZ")
    parser.add_argument("--industry", help="Industry or theme")
    parser.add_argument("--topic", help="Analysis topic")
    parser.add_argument("--market", default="A-share", help="Market, default A-share")
    parser.add_argument("--out", help="Output run directory")
    return parser.parse_args(argv)


def main() -> int:
    run(parse_args())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
