#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REQUIRED_FILES = [
    "analysis.md",
    "evidence_ledger.json",
    "competitor_map.md",
    "industry_structure.md",
    "substitution_risk.md",
    "factor_score.xlsx",
    "catalyst_calendar.md",
    "tracking_plan.md",
    "export_input.json",
    "analysis-pack.json",
    "deep-analysis-prompt.md",
    "handoff_command.txt",
]

REQUIRED_SHEETS = [
    "Factor Score",
    "Evidence Ledger",
    "Sources",
    "Competitors",
    "Substitution Risk",
    "Catalysts",
    "Tracking",
    "Output Handoff",
]


def validate_run(run_dir: Path) -> dict[str, Any]:
    errors: list[str] = []
    warnings: list[str] = []
    if not run_dir.exists():
        return {"status": "failed", "errors": [f"run directory not found: {run_dir}"], "warnings": []}

    for name in REQUIRED_FILES:
        path = run_dir / name
        if not path.exists():
            errors.append(f"missing required file: {name}")
        elif path.stat().st_size == 0:
            errors.append(f"empty required file: {name}")

    export_input_path = run_dir / "export_input.json"
    export_input: dict[str, Any] = {}
    if export_input_path.exists():
        export_input = json.loads(export_input_path.read_text(encoding="utf-8"))
        for key in ["company", "industry", "topic", "sources", "financial_model", "analysis_context"]:
            if key not in export_input:
                errors.append(f"export_input missing key: {key}")
        context = export_input.get("analysis_context", {})
        for key in ["claims", "factor_scores", "competitors", "substitution_risks", "catalysts", "tracking_indicators"]:
            if key not in context:
                errors.append(f"analysis_context missing key: {key}")
    else:
        errors.append("export_input.json not found")

    ledger_path = run_dir / "evidence_ledger.json"
    if ledger_path.exists():
        ledger = json.loads(ledger_path.read_text(encoding="utf-8"))
        levels = ledger.get("levels", {})
        if not {"A", "B", "C", "D"}.issubset(levels):
            errors.append("evidence_ledger missing A/B/C/D level definitions")
        if not isinstance(ledger.get("claims", []), list):
            errors.append("evidence_ledger claims must be a list")

    workbook_path = run_dir / "factor_score.xlsx"
    workbook_stats: dict[str, Any] = {}
    if workbook_path.exists():
        wb = load_workbook(workbook_path, data_only=False)
        missing = [name for name in REQUIRED_SHEETS if name not in wb.sheetnames]
        if missing:
            errors.append("missing factor workbook sheets: " + ", ".join(missing))
        ws = wb["Factor Score"] if "Factor Score" in wb.sheetnames else None
        if ws and ws.max_row < 10:
            errors.append("Factor Score sheet has too few rows")
        formulas = 0
        for ws2 in wb.worksheets:
            for row in ws2.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formulas += 1
        if formulas < 5:
            warnings.append(f"formula count lower than expected: {formulas}")
        workbook_stats = {"sheets": len(wb.sheetnames), "formulas": formulas}

    analysis_path = run_dir / "analysis.md"
    if analysis_path.exists():
        text = analysis_path.read_text(encoding="utf-8")
        for phrase in ["证据等级", "A股主题因子", "竞品分析", "未来行业格局", "替代风险", "下游交付"]:
            if phrase not in text:
                errors.append(f"analysis.md missing section phrase: {phrase}")

    return {
        "status": "passed" if not errors else "failed",
        "run_dir": str(run_dir),
        "errors": errors,
        "warnings": warnings,
        "workbook": workbook_stats,
        "export_company": export_input.get("company"),
    }


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description="Validate a local-analysis-workflow run directory")
    parser.add_argument("run_dir")
    args = parser.parse_args(argv)
    result = validate_run(Path(args.run_dir))
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
